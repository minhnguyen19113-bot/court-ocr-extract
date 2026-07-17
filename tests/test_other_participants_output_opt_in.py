from __future__ import annotations

import json

from openpyxl import load_workbook

from court_ocr_extract.ocr_backends.base import OCRPage, OCRResult
from court_ocr_extract.ocr_cache import OCRCacheRecord
from court_ocr_extract.pre_content_ab import run_pre_content_ab_test
from court_ocr_extract.settings import PipelineSettings


def test_other_participant_output_is_off_by_default_and_opt_in(tmp_path) -> None:
    default_dir = tmp_path / "default"
    run_pre_content_ab_test(
        [_record()],
        output_dir=default_dir,
        settings=PipelineSettings(),
        strategies=["rule_anchor_only"],
    )

    default_workbook = load_workbook(default_dir / "compare_summary.xlsx", read_only=True)
    assert "NGUOI_THAM_GIA_KHAC" not in default_workbook.sheetnames
    assert "PARTICIPANTS" in default_workbook.sheetnames
    assert default_workbook["PARTICIPANTS"].max_row >= 2
    output = json.loads(
        (default_dir / "cases" / "synthetic_case" / "rule_anchor_only_output.json")
        .read_text(encoding="utf-8")
    )
    assert output["participants"]
    default_html = (default_dir / "index.html").read_text(encoding="utf-8")
    assert "<h2>NGƯỜI THAM GIA KHÁC</h2>" not in default_html

    opt_in_dir = tmp_path / "opt_in"
    run_pre_content_ab_test(
        [_record()],
        output_dir=opt_in_dir,
        settings=PipelineSettings(),
        strategies=["rule_anchor_only"],
        include_other_participants_output=True,
    )
    opt_in_workbook = load_workbook(
        opt_in_dir / "compare_summary.xlsx",
        read_only=True,
    )
    assert "NGUOI_THAM_GIA_KHAC" in opt_in_workbook.sheetnames
    assert "<h2>NGƯỜI THAM GIA KHÁC</h2>" in (
        opt_in_dir / "index.html"
    ).read_text(encoding="utf-8")


def _record() -> OCRCacheRecord:
    texts = [
        "Bản án số: 901/2099/HS-ST",
        "Ngày: 01/02/2099",
        "Đối với bị cáo:",
        "1. Person Synthetic Alpha, sinh năm 1990",
        "Nơi ở hiện nay: Zone Synthetic Alpha",
        "Người giám hộ: Ông Person Synthetic Guardian",
        "NỘI DUNG VỤ ÁN",
    ]
    lines = [
        {"line_id": f"p001_l{index:04d}", "page_number": 1, "text": text}
        for index, text in enumerate(texts, start=1)
    ]
    return OCRCacheRecord(
        "synthetic_case",
        1,
        "synthetic_hash",
        OCRResult(
            backend="synthetic",
            status="success",
            pages_processed=1,
            marker_found=True,
            marker_page=1,
            text="\n".join(texts),
            pages=[OCRPage(page_index=1, lines=lines)],
        ),
    )
