from __future__ import annotations

import json

from openpyxl import load_workbook

from court_ocr_extract.ocr_backends.base import OCRPage, OCRResult
from court_ocr_extract.ocr_cache import OCRCacheRecord
from court_ocr_extract.pre_content_ab import DEFAULT_STRATEGIES, STRATEGIES, run_pre_content_ab_test
from court_ocr_extract.settings import PipelineSettings


def test_compare_runner_uses_anchor_defaults_and_bounded_per_block_llm(tmp_path) -> None:
    requests = []

    def fake_llm(prompt, text):
        request = json.loads(text)
        requests.append(request)
        assert len(text) <= 6000
        assert set(("block_id", "type", "line_ids", "raw_lines", "expected_schema")) <= set(request)
        if request["type"] == "defendant":
            return {
                "full_name": "Người Synthetic A",
                "occupation": "Kiểm thử",
                "permanent_address": "Vùng Synthetic A",
                "detention_status": "Cấm đi khỏi nơi cư trú",
            }
        return {
            "role": request["role_hint"],
            "full_name": "Người Synthetic P",
            "address": "Vùng Synthetic P",
            "relationship_or_note": "Vai trò Synthetic",
        }

    summary = run_pre_content_ab_test(
        [_record()],
        output_dir=tmp_path,
        settings=PipelineSettings(local_llm_max_output_tokens=1024, local_llm_max_input_chars=22000),
        llm_callable=fake_llm,
    )

    assert summary["strategies"] == list(DEFAULT_STRATEGIES)
    assert requests
    assert {request["type"] for request in requests} == {"defendant", "participant"}
    assert "hybrid_rule_llm" in STRATEGIES
    assert "legacy_hybrid_rule_llm" in STRATEGIES
    case_dir = tmp_path / "cases" / "synthetic"
    assert (case_dir / "rule_anchor_only_output.json").exists()
    repaired = json.loads((case_dir / "rule_then_llm_per_block_output.json").read_text(encoding="utf-8"))
    assert repaired["metadata"]["judgment_number"] == "01/2025/HS-ST"
    assert repaired["defendants"][0]["occupation"] == "Kiểm thử"
    assert repaired["participants"][0]["relationship_or_note"] == "Vai trò Synthetic"
    assert all(status["max_output_tokens"] <= 512 for status in repaired["llm_status"])
    workbook = load_workbook(tmp_path / "compare_summary.xlsx", read_only=True)
    assert {
        "ANCHOR_BLOCKS", "ANCHOR_WARNINGS", "CASES", "DEFENDANTS", "PARTICIPANTS",
        "TRIAL_PANEL", "LLM_STATUS", "FIELD_LONG", "EVIDENCE_LINES", "RAW_JSON",
    } <= set(workbook.sheetnames)
    html = (case_dir / "review.html").read_text(encoding="utf-8")
    assert "Anchor segmentation" in html
    assert "defendant_001" in html
    assert "participant_001" in html


def test_rule_anchor_only_does_not_call_llm(tmp_path) -> None:
    def forbidden_llm(prompt, text):
        raise AssertionError("Rule-only strategy must not call LLM")

    summary = run_pre_content_ab_test(
        [_record()],
        output_dir=tmp_path,
        settings=PipelineSettings(),
        strategies=["rule_anchor_only"],
        llm_callable=forbidden_llm,
    )

    assert summary["totals"]["llm_chunk_count"] == 0


def test_llm_per_block_all_failed_is_not_a_valid_result(tmp_path) -> None:
    def failed_llm(prompt, text):
        raise RuntimeError("synthetic block failure")

    run_pre_content_ab_test(
        [_record()],
        output_dir=tmp_path,
        settings=PipelineSettings(),
        strategies=["llm_per_block"],
        llm_callable=failed_llm,
    )

    payload = json.loads(
        (tmp_path / "cases" / "synthetic" / "llm_per_block_output.json").read_text(
            encoding="utf-8"
        )
    )
    assert payload["status"] == "llm_per_block_failed"
    assert payload["result_valid"] is False
    assert payload["llm_actually_called"] is True
    assert payload["warnings"]


def _record() -> OCRCacheRecord:
    texts = [
        "TÒA ÁN NHÂN DÂN SYNTHETIC",
        "Bản án số: 01/2025/HS-ST Ngày: 01/7/2025",
        "Thẩm phán - Chủ tọa phiên tòa: Thẩm Phán Synthetic",
        "Đối với bị cáo:",
        "Người Synthetic A, sinh năm 1990",
        "Bị hại:",
        "Người Synthetic P",
        "NỘI DUNG VỤ ÁN",
    ]
    lines = [
        {"line_id": f"p001_l{index:04d}", "page_number": 1, "text": text}
        for index, text in enumerate(texts, start=1)
    ]
    return OCRCacheRecord(
        "synthetic",
        1,
        "synthetic",
        OCRResult(
            backend="synthetic",
            status="success",
            pages_processed=1,
            marker_found=True,
            marker_page=1,
            text="\n".join(texts),
            pages=[OCRPage(page_index=1, lines=lines)],
        ),
    )
