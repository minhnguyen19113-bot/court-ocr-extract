from openpyxl import load_workbook

from court_ocr_extract.excel_writer import (
    EXCEL_HEADERS,
    write_excel,
    write_excel_from_results,
)
from court_ocr_extract.models import CaseInfo, ExtractionResult, Participant
from court_ocr_extract.other_participants_builder import (
    OTHER_PARTICIPANTS_SHEET_NAME,
)


def test_canonical_writer_creates_draft_workbook_contract(tmp_path):
    output_path = tmp_path / "synthetic_draft.xlsx"
    drafts = [
        {
            "case_id": "synthetic_case_001",
            "status": "success",
            "marker_found": True,
            "payload": {
                "case": {
                    "case_type": "Loại án synthetic",
                    "filing_number": "SYNTHETIC-001",
                },
                "participants": [
                    {
                        "procedural_role": "Bị hại",
                        "full_name": "Người kiểm tra synthetic",
                    }
                ],
            },
        }
    ]

    result_path = write_excel(drafts, output_path)
    workbook = load_workbook(result_path, read_only=True)

    assert workbook.sheetnames == ["FINAL_EXCEL", "RUN_SUMMARY"]
    assert OTHER_PARTICIPANTS_SHEET_NAME not in workbook.sheetnames
    assert [cell.value for cell in workbook["FINAL_EXCEL"][1]] == EXCEL_HEADERS
    assert workbook["FINAL_EXCEL"].max_row == 2


def test_canonical_writer_creates_typed_result_workbook_contract(tmp_path):
    output_path = tmp_path / "synthetic_typed.xlsx"
    result = ExtractionResult(
        case_info=CaseInfo(
            loai_an="Loại án synthetic",
            so_thu_ly="SYNTHETIC-002",
        ),
        participants=[
            Participant(
                tu_cach_to_tung="Bị hại",
                ho_ten="Người kiểm tra synthetic",
            ),
            Participant(
                tu_cach_to_tung="Người giám hộ",
                ho_ten="Người hỗ trợ synthetic",
            ),
        ],
    )

    result_path = write_excel_from_results(
        [result],
        output_path,
        include_other_participants_output=True,
    )
    workbook = load_workbook(result_path, read_only=True)

    assert workbook.sheetnames == ["FINAL_EXCEL", OTHER_PARTICIPANTS_SHEET_NAME]
    assert [cell.value for cell in workbook["FINAL_EXCEL"][1]] == EXCEL_HEADERS
    assert workbook["FINAL_EXCEL"].max_row == 2
    assert _column_values(workbook["FINAL_EXCEL"], "TƯ CÁCH TỐ TỤNG") == [
        "Bị hại"
    ]
    assert _column_values(
        workbook[OTHER_PARTICIPANTS_SHEET_NAME],
        "TƯ CÁCH TỐ TỤNG",
    ) == ["Người giám hộ"]


def _column_values(sheet, header: str) -> list[object]:
    headers = next(sheet.iter_rows(max_row=1, values_only=True))
    header_index = {value: index for index, value in enumerate(headers)}
    return [
        row[header_index[header]]
        for row in sheet.iter_rows(min_row=2, values_only=True)
    ]
