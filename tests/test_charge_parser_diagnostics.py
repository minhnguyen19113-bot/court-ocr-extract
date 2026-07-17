from __future__ import annotations

import json

from openpyxl import load_workbook

from court_ocr_extract.decision_tail import DecisionTailRecord
from court_ocr_extract.ocr_backends.base import OCRPage, OCRResult
from court_ocr_extract.ocr_cache import OCRCacheRecord
from court_ocr_extract.pre_content_ab import run_pre_content_ab_test
from court_ocr_extract.settings import PipelineSettings


def test_heading_without_verdict_candidate_has_explicit_diagnostics(tmp_path) -> None:
    tail = DecisionTailRecord(
        case_id="synthetic_case",
        source_index=1,
        pdf_hash="synthetic_hash",
        backend="synthetic",
        pages_total=9,
        scanned_page_numbers=[9],
        scan_batches=[[9]],
        heading_found=True,
        heading_page=9,
        heading_line_id="p009_l0001",
        heading_text="QUYẾT ĐỊNH",
        text="QUYẾT ĐỊNH\nÁp dụng quy định synthetic.",
        lines=[
            {"line_id": "p009_l0001", "page_number": 9, "text": "QUYẾT ĐỊNH"},
            {
                "line_id": "p009_l0002",
                "page_number": 9,
                "text": "Áp dụng quy định synthetic.",
            },
        ],
    )
    run_pre_content_ab_test(
        [_front_record()],
        output_dir=tmp_path,
        settings=PipelineSettings(),
        strategies=["rule_anchor_only"],
        decision_tail_records={"synthetic_case": tail},
    )

    output = json.loads(
        (tmp_path / "cases" / "synthetic_case" / "rule_anchor_only_output.json")
        .read_text(encoding="utf-8")
    )
    assert output["decision_tail_status"] == "charge_not_found"
    assert output["decision_heading_found"] is True
    assert output["decision_heading_page"] == 9
    assert output["decision_tail_line_count"] == 2
    assert output["verdict_candidate_count"] == 0
    assert output["parsed_charge_count"] == 0
    assert output["mapped_defendant_count"] == 0
    assert output["unmapped_verdict_count"] == 0
    assert "decision_heading_found_but_no_verdict_candidate" in output["warnings"]

    workbook = load_workbook(tmp_path / "compare_summary.xlsx", read_only=True)
    headers = [cell.value for cell in next(workbook["CASES"].iter_rows(max_row=1))]
    values = [cell.value for cell in next(workbook["CASES"].iter_rows(min_row=2, max_row=2))]
    case_row = dict(zip(headers, values))
    assert case_row["verdict_candidate_count"] == 0
    html = (tmp_path / "index.html").read_text(encoding="utf-8")
    assert "heading found: True" in html
    assert "verdict candidates: 0" in html


def _front_record() -> OCRCacheRecord:
    texts = [
        "Bản án số: 901/2099/HS-ST",
        "Đối với bị cáo:",
        "1. Person Synthetic Alpha, sinh năm 1990",
        "NỘI DUNG VỤ ÁN",
    ]
    lines = [
        {"line_id": f"p001_l{index:04d}", "page_number": 1, "text": text}
        for index, text in enumerate(texts, start=1)
    ]
    return OCRCacheRecord(
        "synthetic_case",
        1,
        "synthetic_hash",
        OCRResult(
            backend="synthetic",
            status="success",
            pages_processed=1,
            marker_found=True,
            marker_page=1,
            text="\n".join(texts),
            pages=[OCRPage(page_index=1, lines=lines)],
        ),
    )
