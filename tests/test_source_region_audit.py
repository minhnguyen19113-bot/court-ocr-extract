from __future__ import annotations

from openpyxl import load_workbook

from court_ocr_extract.decision_tail import DecisionTailRecord
from court_ocr_extract.ocr_backends.base import OCRPage, OCRResult
from court_ocr_extract.ocr_cache import OCRCacheRecord
from court_ocr_extract.pre_content_ab import run_pre_content_ab_test
from court_ocr_extract.settings import PipelineSettings
from court_ocr_extract.source_region_policy import DECISION_TAIL, MIDDLE_EXCLUDED


def test_workbook_and_html_expose_source_and_entity_charge_audit(tmp_path) -> None:
    record = _front_record()
    tail = DecisionTailRecord(
        case_id=record.case_id,
        source_index=1,
        pdf_hash="synthetic_hash",
        backend="synthetic",
        pages_total=10,
        scanned_page_numbers=[9, 10],
        scan_batches=[[9, 10]],
        heading_found=True,
        heading_page=9,
        heading_line_id="p009_l0001",
        heading_text="QUYẾT ĐỊNH",
        text="",
        lines=[
            {
                "line_id": "p005_l0001",
                "page_number": 5,
                "source_region": MIDDLE_EXCLUDED,
                "text": (
                    "Tuyên bị cáo Person Synthetic Alpha phạm tội "
                    '"Charge Synthetic Misleading".'
                ),
            },
            {
                "line_id": "p009_l0001",
                "page_number": 9,
                "source_region": DECISION_TAIL,
                "text": "QUYẾT ĐỊNH",
            },
            {
                "line_id": "p009_l0002",
                "page_number": 9,
                "source_region": DECISION_TAIL,
                "text": (
                    "Tuyên các bị cáo Person Synthetic Alpha, Person Synthetic Beta "
                    'phạm tội "Charge Synthetic A".'
                ),
            },
            {
                "line_id": "p010_l0001",
                "page_number": 10,
                "source_region": DECISION_TAIL,
                "text": (
                    "Tuyên bị cáo Person Synthetic Gamma phạm tội "
                    '"Charge Synthetic B".'
                ),
            },
        ],
    )
    run_pre_content_ab_test(
        [record],
        output_dir=tmp_path,
        settings=PipelineSettings(),
        strategies=["rule_anchor_only"],
        decision_tail_records={record.case_id: tail},
    )

    workbook = load_workbook(tmp_path / "compare_summary.xlsx", read_only=True)
    for sheet_name in (
        "CHARGES",
        "DEFENDANT_CHARGES",
        "SOURCE_REGION_AUDIT",
        "CHARGE_WARNINGS",
    ):
        assert sheet_name in workbook.sheetnames
    audit_regions = {
        row[2]
        for row in workbook["SOURCE_REGION_AUDIT"].iter_rows(
            min_row=2,
            values_only=True,
        )
    }
    assert audit_regions <= {"front_pre_content", "decision_tail"}

    html = (tmp_path / "cases" / record.case_id / "review.html").read_text(
        encoding="utf-8"
    )
    assert html.index("FINAL EXCEL PREVIEW") < html.index("NGƯỜI THAM GIA KHÁC")
    assert html.index("NGƯỜI THAM GIA KHÁC") < html.index("CHARGE SUMMARY")
    assert html.index("CHARGE SUMMARY") < html.index("DEFENDANT → CHARGE MAP")
    assert html.index("DEFENDANT → CHARGE MAP") < html.index("SOURCE REGION AUDIT")
    assert "Charge Synthetic Misleading" not in html


def _front_record() -> OCRCacheRecord:
    texts = [
        "Bản án số: 950/2099/HS-ST",
        "Đối với các bị cáo:",
        "1. Person Synthetic Alpha, sinh năm 1990",
        "Nơi ở hiện nay: Region Synthetic Alpha",
        "2. Person Synthetic Beta, sinh năm 1991",
        "Nơi ở hiện nay: Region Synthetic Beta",
        "3. Person Synthetic Gamma, sinh năm 1992",
        "Nơi ở hiện nay: Region Synthetic Gamma",
        "Bị hại: Person Synthetic Alpha",
        "NỘI DUNG VỤ ÁN",
    ]
    lines = [
        {"line_id": f"p001_l{index:04d}", "page_number": 1, "text": text}
        for index, text in enumerate(texts, start=1)
    ]
    return OCRCacheRecord(
        "synthetic_case",
        1,
        "synthetic_hash",
        OCRResult(
            backend="synthetic",
            status="success",
            pages_processed=1,
            marker_found=True,
            marker_page=1,
            text="\n".join(texts),
            pages=[OCRPage(page_index=1, lines=lines)],
        ),
    )

