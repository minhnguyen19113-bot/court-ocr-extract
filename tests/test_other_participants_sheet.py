from __future__ import annotations

from openpyxl import load_workbook

from court_ocr_extract.final_excel_schema import FINAL_EXCEL_COLUMNS
from court_ocr_extract.ocr_backends.base import OCRPage, OCRResult
from court_ocr_extract.ocr_cache import OCRCacheRecord
from court_ocr_extract.other_participants_builder import (
    OTHER_PARTICIPANT_COLUMNS,
    OTHER_PARTICIPANTS_SHEET_NAME,
)
from court_ocr_extract.pre_content_ab import run_pre_content_ab_test
from court_ocr_extract.settings import PipelineSettings


def test_other_participants_are_user_facing_but_not_in_final_excel(tmp_path) -> None:
    texts = [
        "Bản án số: 909/2099/HS-ST",
        "Thẩm phán - Chủ tọa phiên tòa: Chủ Tọa Synthetic",
        "Đối với các bị cáo:",
        "1. Người Synthetic A, sinh năm 1990",
        "Nơi ở hiện nay: Vùng Synthetic A",
        "2. Người Synthetic B, sinh năm 1991",
        "Nơi ở hiện nay: Vùng Synthetic B",
        "Bị hại: Người Synthetic C",
        "Người giám hộ: Ông Synthetic D",
        "Người bảo vệ quyền và lợi ích hợp pháp của bị hại: Bà Synthetic E",
        "Người bào chữa cho bị cáo: Ông Synthetic F",
        "NỘI DUNG VỤ ÁN",
    ]
    run_pre_content_ab_test(
        [_record(texts)],
        output_dir=tmp_path,
        settings=PipelineSettings(),
        strategies=["rule_anchor_only"],
        include_other_participants_output=True,
    )

    workbook = load_workbook(tmp_path / "compare_summary.xlsx", read_only=True)
    assert workbook.sheetnames[:2] == ["FINAL_EXCEL", OTHER_PARTICIPANTS_SHEET_NAME]
    assert workbook["FINAL_EXCEL"].max_row == 4
    assert [cell.value for cell in workbook[OTHER_PARTICIPANTS_SHEET_NAME][1]] == (
        OTHER_PARTICIPANT_COLUMNS
    )
    assert workbook[OTHER_PARTICIPANTS_SHEET_NAME].max_row == 4
    final_header_index = {
        value: index for index, value in enumerate(FINAL_EXCEL_COLUMNS)
    }
    final_roles = {
        row[final_header_index["TƯ CÁCH TỐ TỤNG"]]
        for row in workbook["FINAL_EXCEL"].iter_rows(min_row=2, values_only=True)
    }
    other_header_index = {
        value: index for index, value in enumerate(OTHER_PARTICIPANT_COLUMNS)
    }
    other_roles = {
        row[other_header_index["TƯ CÁCH TỐ TỤNG"]]
        for row in workbook[OTHER_PARTICIPANTS_SHEET_NAME].iter_rows(
            min_row=2, values_only=True
        )
    }
    assert final_roles == {"Bị cáo", "Bị hại"}
    assert other_roles == {
        "Người giám hộ",
        "Người bảo vệ quyền và lợi ích hợp pháp của bị hại",
        "Người bào chữa cho bị cáo",
    }

    html = (tmp_path / "cases" / "synthetic_case" / "review.html").read_text(
        encoding="utf-8"
    )
    assert html.index("FINAL EXCEL PREVIEW") < html.index("NGƯỜI THAM GIA KHÁC")
    assert html.index("NGƯỜI THAM GIA KHÁC") < html.index("Anchor segmentation")


def _record(texts: list[str]) -> OCRCacheRecord:
    lines = [
        {"line_id": f"p001_l{index:04d}", "page_number": 1, "text": text}
        for index, text in enumerate(texts, start=1)
    ]
    return OCRCacheRecord(
        "synthetic_case",
        1,
        "synthetic_hash",
        OCRResult(
            backend="synthetic",
            status="success",
            pages_processed=1,
            marker_found=True,
            marker_page=1,
            text="\n".join(texts),
            pages=[OCRPage(page_index=1, lines=lines)],
        ),
    )
