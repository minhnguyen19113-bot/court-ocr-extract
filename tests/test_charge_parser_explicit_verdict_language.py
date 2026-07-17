from __future__ import annotations

from openpyxl import load_workbook

from court_ocr_extract.charge_parser import parse_explicit_decision_charges
from court_ocr_extract.decision_tail import DecisionTailRecord
from court_ocr_extract.ocr_backends.base import OCRPage, OCRResult
from court_ocr_extract.ocr_cache import OCRCacheRecord
from court_ocr_extract.pre_content_ab import run_pre_content_ab_test
from court_ocr_extract.settings import PipelineSettings
from court_ocr_extract.source_region_policy import DECISION_TAIL


def test_charge_parser_accepts_only_explicit_verdict_language_and_quote_variants() -> None:
    lines = [
        {
            "line_id": "p009_l0001",
            "source_region": DECISION_TAIL,
            "text": "Tuyên bị cáo Người Synthetic A phạm tội “Tội Synthetic Alpha”",
        },
        {
            "line_id": "p009_l0002",
            "source_region": DECISION_TAIL,
            "text": "Xử phạt bị cáo Người Synthetic B về tội 'Tội Synthetic Beta'",
        },
        {
            "line_id": "p009_l0003",
            "source_region": DECISION_TAIL,
            "text": "Áp dụng Điều synthetic cho mô tả hành vi synthetic",
        },
    ]
    defendants = [
        {"entity_id": "defendant_a", "full_name": "Người Synthetic A"},
        {"entity_id": "defendant_b", "full_name": "Người Synthetic B"},
    ]

    output = parse_explicit_decision_charges(lines, defendants=defendants)

    assert output["case_charges"] == ["Tội Synthetic Alpha", "Tội Synthetic Beta"]
    assert output["defendant_charge_map"] == {
        "defendant_a": ["Tội Synthetic Alpha"],
        "defendant_b": ["Tội Synthetic Beta"],
    }
    assert len(output["charge_evidence"]) == 2
    assert all(
        item["source_region"] == DECISION_TAIL
        for item in output["charge_evidence"]
    )
    assert output["warnings"] == []


def test_charge_parser_does_not_infer_from_law_or_behavior_only() -> None:
    output = parse_explicit_decision_charges(
        "Áp dụng Điều synthetic. Mô tả hành vi synthetic không có câu tuyên tội."
    )

    assert output == {
        "case_charges": [],
        "defendant_charge_map": {},
        "charge_evidence": [],
        "warnings": [],
    }


def test_tail_cache_charges_are_repeated_in_final_excel(tmp_path) -> None:
    record = _pre_content_record()
    tail = DecisionTailRecord(
        case_id=record.case_id,
        source_index=1,
        pdf_hash="synthetic_hash",
        backend="synthetic",
        pages_total=9,
        scanned_page_numbers=[9],
        scan_batches=[[9]],
        heading_found=True,
        heading_page=9,
        heading_line_id="p009_l0001",
        heading_text="QUYẾT ĐỊNH",
        text="",
        lines=[
            {"line_id": "p009_l0001", "page_number": 9, "text": "QUYẾT ĐỊNH"},
            {
                "line_id": "p009_l0002",
                "page_number": 9,
                "text": (
                    "Tuyên các bị cáo Người Synthetic A và Người Synthetic B "
                    "phạm tội \"Tội Synthetic Alpha\""
                ),
            },
        ],
    )
    run_pre_content_ab_test(
        [record],
        output_dir=tmp_path,
        settings=PipelineSettings(),
        strategies=["rule_anchor_only"],
        decision_tail_records={record.case_id: tail},
    )

    workbook = load_workbook(tmp_path / "compare_summary.xlsx", read_only=True)
    relationships = {
        row[3]
        for row in workbook["FINAL_EXCEL"].iter_rows(min_row=2, values_only=True)
    }
    assert relationships == {"Tội Synthetic Alpha"}
    assert workbook["CHARGES"].max_row == 2
    assert workbook["DEFENDANT_CHARGES"].max_row == 3


def _pre_content_record() -> OCRCacheRecord:
    texts = [
        "Bản án số: 911/2099/HS-ST",
        "Đối với các bị cáo:",
        "1. Người Synthetic A, sinh năm 1990",
        "Nơi ở hiện nay: Vùng Synthetic A",
        "2. Người Synthetic B, sinh năm 1991",
        "Nơi ở hiện nay: Vùng Synthetic B",
        "NỘI DUNG VỤ ÁN",
    ]
    lines = [
        {"line_id": f"p001_l{index:04d}", "page_number": 1, "text": text}
        for index, text in enumerate(texts, start=1)
    ]
    return OCRCacheRecord(
        "synthetic_case",
        1,
        "synthetic_hash",
        OCRResult(
            backend="synthetic",
            status="success",
            pages_processed=1,
            marker_found=True,
            marker_page=1,
            text="\n".join(texts),
            pages=[OCRPage(page_index=1, lines=lines)],
        ),
    )
