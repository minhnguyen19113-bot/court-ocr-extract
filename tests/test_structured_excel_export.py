from __future__ import annotations

import json

from openpyxl import load_workbook

from court_ocr_extract.pre_content_ab import run_pre_content_ab_test
from court_ocr_extract.settings import PipelineSettings
from tests.test_compare_requires_llm import _record


def test_compare_workbook_has_structured_review_sheets_without_json_blobs(tmp_path) -> None:
    def fake_llm(prompt, text):
        target = json.loads(text)["target"]
        if target == "document_metadata":
            return {"metadata": {"court_name": "Synthetic Court", "judgment_number": "01/SYNTHETIC"}}
        if target == "trial_panel":
            return {"trial_panel": {"presiding_judge": "Synthetic Judge"}}
        if target == "defendants":
            return {"defendants": [{"full_name": "Synthetic Person", "evidence_line_ids": ["p001_l0001"]}]}
        return {"participants": [{"role": "synthetic role", "full_name": "Synthetic Participant"}]}

    run_pre_content_ab_test(
        [_record()], output_dir=tmp_path, settings=PipelineSettings(), llm_callable=fake_llm
    )
    workbook = load_workbook(tmp_path / "compare_summary.xlsx", read_only=True)
    required = {"CASES", "DEFENDANTS", "PARTICIPANTS", "TRIAL_PANEL", "LLM_STATUS", "FIELD_LONG", "EVIDENCE_LINES", "RAW_JSON"}
    assert required <= set(workbook.sheetnames)
    assert [cell.value for cell in next(workbook["CASES"].iter_rows())] == list(__import__("court_ocr_extract.pre_content_ab", fromlist=["CASES_HEADERS"]).CASES_HEADERS)
    for sheet_name in ("CASES", "DEFENDANTS", "PARTICIPANTS", "TRIAL_PANEL"):
        assert workbook[sheet_name].max_row >= 2
        for row in workbook[sheet_name].iter_rows(min_row=2, values_only=True):
            assert not any(isinstance(value, str) and value.lstrip().startswith(("{", "[")) for value in row)
