from __future__ import annotations

from openpyxl import load_workbook

from court_ocr_extract.decision_tail import DecisionTailRecord
from court_ocr_extract.ocr_backends.base import OCRPage, OCRResult
from court_ocr_extract.ocr_cache import OCRCacheRecord
from court_ocr_extract.pre_content_ab import run_pre_content_ab_test
from court_ocr_extract.settings import PipelineSettings
from court_ocr_extract.source_region_policy import DECISION_TAIL


def test_sentence_workbook_audit_sheets_and_final_column(tmp_path) -> None:
    record = _front_record()
    tail = DecisionTailRecord(
        case_id=record.case_id,
        source_index=1,
        pdf_hash="synthetic_hash",
        backend="synthetic",
        pages_total=9,
        scanned_page_numbers=[9],
        scan_batches=[[9]],
        heading_found=True,
        heading_page=9,
        heading_line_id="p009_l0001",
        heading_text="QUYẾT ĐỊNH",
        text="",
        lines=[
            {
                "line_id": "p009_l0001",
                "page_number": 9,
                "source_region": DECISION_TAIL,
                "text": "QUYẾT ĐỊNH",
            },
            {
                "line_id": "p009_l0002",
                "page_number": 9,
                "source_region": DECISION_TAIL,
                "text": (
                    "Xử phạt bị cáo Person Synthetic Alpha 02 năm tù "
                    "về tội “Charge Synthetic”."
                ),
            },
            {
                "line_id": "p009_l0003",
                "page_number": 9,
                "source_region": DECISION_TAIL,
                "text": "Thời hạn tù tính từ ngày chưa xác định.",
            },
        ],
    )
    run_pre_content_ab_test(
        [record],
        output_dir=tmp_path,
        settings=PipelineSettings(),
        strategies=["rule_anchor_only"],
        decision_tail_records={record.case_id: tail},
    )

    workbook = load_workbook(tmp_path / "compare_summary.xlsx", read_only=True)
    assert {
        "DEFENDANT_SENTENCES",
        "SENTENCE_EVIDENCE",
        "SENTENCE_WARNINGS",
    } <= set(workbook.sheetnames)
    final_sheet = workbook["FINAL_EXCEL"]
    headers = list(next(final_sheet.iter_rows(max_row=1, values_only=True)))
    sentence_index = headers.index("HÌNH PHẠT")
    row = next(final_sheet.iter_rows(min_row=2, values_only=True))
    assert row[sentence_index] == "2 năm tù"
    assert workbook["DEFENDANT_SENTENCES"].max_row == 2
    assert workbook["SENTENCE_EVIDENCE"].max_row == 2
    evidence_headers = list(
        next(workbook["SENTENCE_EVIDENCE"].iter_rows(max_row=1, values_only=True))
    )
    assert evidence_headers == [
        "case_id", "evidence_type", "defendant_entity_ids", "defendant_names",
        "raw_text", "page_number", "line_ids", "source_region", "match_method",
        "confidence", "warnings",
    ]
    warning_rows = list(workbook["SENTENCE_WARNINGS"].iter_rows(values_only=True))
    warning_headers = list(warning_rows[0])
    warning = dict(zip(warning_headers, warning_rows[1], strict=True))
    assert warning["warning"] == "sentence_start_anchor_unparsed"
    assert warning["scope"] == "entity"
    assert warning["defendant_entity_id"] == "defendant_001"
    assert warning["page_number"] == 9
    assert warning["line_ids"] == "p009_l0003"
    assert warning["raw_text"] == "Thời hạn tù tính từ ngày chưa xác định."


def _front_record() -> OCRCacheRecord:
    texts = [
        "Bản án số: 901/2099/HS-ST",
        "Đối với bị cáo:",
        "1. Person Synthetic Alpha, sinh năm 1990",
        "Nơi ở: Vùng Synthetic Current",
        "NỘI DUNG VỤ ÁN",
    ]
    lines = [
        {"line_id": f"p001_l{index:04d}", "page_number": 1, "text": text}
        for index, text in enumerate(texts, start=1)
    ]
    return OCRCacheRecord(
        "synthetic_sentence_case",
        1,
        "synthetic_hash",
        OCRResult(
            backend="synthetic",
            status="success",
            pages_processed=1,
            marker_found=True,
            marker_page=1,
            text="\n".join(texts),
            pages=[OCRPage(page_index=1, lines=lines)],
        ),
    )
