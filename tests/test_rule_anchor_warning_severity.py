from __future__ import annotations

from openpyxl import load_workbook

from court_ocr_extract.ocr_backends.base import OCRPage, OCRResult
from court_ocr_extract.ocr_cache import OCRCacheRecord
from court_ocr_extract.pre_content_ab import run_pre_content_ab_test
from court_ocr_extract.settings import PipelineSettings


def test_standalone_page_warning_is_exported_but_not_reviewable(tmp_path) -> None:
    texts = [
        "Bản án số: 01/2025/HS-ST",
        "Thẩm phán - Chủ tọa phiên tòa: Thẩm Phán Synthetic",
        "Đối với bị cáo:",
        "Người Synthetic A, sinh năm 1990",
        "1",
        "NỘI DUNG VỤ ÁN",
    ]

    run_pre_content_ab_test(
        [_record(texts)],
        output_dir=tmp_path,
        settings=PipelineSettings(),
        strategies=["rule_anchor_only"],
    )

    workbook = load_workbook(tmp_path / "compare_summary.xlsx", read_only=True)
    cases = list(workbook["CASES"].iter_rows(values_only=True))
    case_headers = cases[0]
    needs_review_index = case_headers.index("needs_review")
    assert cases[1][needs_review_index] is False
    warnings = list(workbook["ANCHOR_WARNINGS"].iter_rows(min_row=2, values_only=True))
    assert warnings == [("synthetic", "standalone_page_number_removed_from_defendant_blocks")]


def _record(texts: list[str]) -> OCRCacheRecord:
    lines = [
        {"line_id": f"p001_l{index:04d}", "page_number": 1, "text": text}
        for index, text in enumerate(texts, start=1)
    ]
    return OCRCacheRecord(
        "synthetic",
        1,
        "synthetic",
        OCRResult(
            backend="synthetic",
            status="success",
            pages_processed=1,
            marker_found=True,
            marker_page=1,
            text="\n".join(texts),
            pages=[OCRPage(page_index=1, lines=lines)],
        ),
    )
