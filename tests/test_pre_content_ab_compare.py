from __future__ import annotations

import json

from openpyxl import load_workbook

from court_ocr_extract.ocr_backends.base import OCRPage, OCRResult
from court_ocr_extract.ocr_cache import OCRCacheRecord
from court_ocr_extract.pre_content_ab import run_pre_content_ab_test
from court_ocr_extract.settings import PipelineSettings


def test_compare_runner_writes_json_excel_html_and_routes_notice(tmp_path) -> None:
    def fake_llm(prompt, text):
        return {
            "document_type": "judgment_criminal_first_instance",
            "metadata": {"judgment_number": "01/2026/HS-ST"},
            "trial_panel": {}, "defendants": [], "participants": [], "evidence": [],
        }

    records = [
        _record("case_judgment", ["Bản án số: 01/2026/HS-ST", "NHÂN DANH", "NỘI DUNG VỤ ÁN"]),
        _record("case_notice", ["THÔNG BÁO", "SỬA CHỮA BỔ SUNG BẢN ÁN", "NỘI DUNG VỤ ÁN"]),
    ]

    summary = run_pre_content_ab_test(
        records, output_dir=tmp_path / "ab", settings=PipelineSettings(), llm_callable=fake_llm
    )

    assert summary["case_count"] == 2
    assert summary["judgment_benchmark_count"] == 1
    assert summary["correction_notice_count"] == 1
    assert (tmp_path / "ab" / "index.html").exists()
    assert (tmp_path / "ab" / "cases" / "case_judgment" / "compare.json").exists()
    review_html = (tmp_path / "ab" / "cases" / "case_judgment" / "review.html").read_text(encoding="utf-8")
    assert 'id="p001_l0001"' in review_html
    assert 'href="#p001_l0001"' in review_html
    payload = json.loads((tmp_path / "ab" / "compare_summary.json").read_text(encoding="utf-8"))
    assert payload["cases"][0]["segment_stop_found"] is True
    workbook = load_workbook(tmp_path / "ab" / "compare_summary.xlsx", read_only=True)
    assert set(workbook.sheetnames) == {
        "SUMMARY", "CASE_COMPARE", "HYBRID_FIELDS", "LLM_ONLY_FIELDS", "CONFLICTS",
        "MISSING_FIELDS", "NEEDS_REVIEW", "EVIDENCE", "DOC_ROUTER",
    }


def _record(case_id, texts):
    lines = [
        {"line_id": f"p001_l{index:04d}", "page_number": 1, "text": text}
        for index, text in enumerate(texts, start=1)
    ]
    return OCRCacheRecord(
        case_id=case_id, source_index=1, pdf_hash="synthetic",
        result=OCRResult(
            backend="synthetic", status="success", pages_processed=1, marker_found=False,
            marker_page=None, text="\n".join(texts), pages=[OCRPage(page_index=1, lines=lines)],
        ),
    )
