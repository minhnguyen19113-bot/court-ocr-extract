from __future__ import annotations

from openpyxl import load_workbook

from court_ocr_extract.ocr_backends.base import OCRPage, OCRResult
from court_ocr_extract.ocr_cache import OCRCacheRecord
from court_ocr_extract.pre_content_ab import run_pre_content_ab_test
from court_ocr_extract.settings import PipelineSettings


def test_juror_value_with_embedded_newline_becomes_two_rows(tmp_path) -> None:
    texts = [
        "Bản án số: 01/2025/HS-ST",
        "Thành phần Hội đồng xét xử:",
        "Thẩm phán - Chủ tọa phiên tòa: Thẩm Phán Synthetic",
        "Các Hội thẩm nhân dân: Ông Synthetic A\nÔng Synthetic B",
        "Đối với bị cáo:",
        "Người Synthetic D, sinh năm 1990",
        "NỘI DUNG VỤ ÁN",
    ]
    record = _record(texts)

    run_pre_content_ab_test(
        [record],
        output_dir=tmp_path,
        settings=PipelineSettings(),
        strategies=["rule_anchor_only"],
    )

    workbook = load_workbook(tmp_path / "compare_summary.xlsx", read_only=True)
    rows = list(workbook["TRIAL_PANEL"].iter_rows(min_row=2, values_only=True))
    jurors = [row[3] for row in rows if row[2] == "juror"]
    assert jurors == ["Ông Synthetic A", "Ông Synthetic B"]


def _record(texts: list[str]) -> OCRCacheRecord:
    lines = [
        {"line_id": f"p001_l{index:04d}", "page_number": 1, "text": text}
        for index, text in enumerate(texts, start=1)
    ]
    return OCRCacheRecord(
        "synthetic",
        1,
        "synthetic",
        OCRResult(
            backend="synthetic",
            status="success",
            pages_processed=1,
            marker_found=True,
            marker_page=1,
            text="\n".join(texts),
            pages=[OCRPage(page_index=1, lines=lines)],
        ),
    )

