from __future__ import annotations

import json

from openpyxl import load_workbook

from court_ocr_extract.pre_content_ab import run_pre_content_ab_test
from court_ocr_extract.settings import PipelineSettings
from tests.test_compare_requires_llm import _record


def test_html_and_llm_status_sheet_show_runtime_budget_and_chunk_error(tmp_path) -> None:
    def fake_llm(prompt, text):
        request = json.loads(text)
        if request["type"] == "participant":
            raise RuntimeError("synthetic failure")
        return {
            "full_name": "Người Synthetic A",
            "occupation": "Kiểm thử",
            "permanent_address": "Vùng Synthetic A",
            "detention_status": "Cấm đi khỏi nơi cư trú",
        }

    run_pre_content_ab_test(
        [_record()], output_dir=tmp_path, settings=PipelineSettings(), llm_callable=fake_llm,
        llm_preflight={"ok": True, "model": "Qwen/Qwen2.5-3B-Instruct", "base_url": "http://127.0.0.1:8000/v1"},
    )
    html = (tmp_path / "cases" / "synthetic" / "review.html").read_text(encoding="utf-8")
    assert "Trạng thái Local LLM" in html
    assert "Input token ước lượng" in html
    workbook = load_workbook(tmp_path / "compare_summary.xlsx", read_only=True)
    rows = list(workbook["LLM_STATUS"].iter_rows(values_only=True))
    headers = rows[0]
    model_index = headers.index("model")
    budget_index = headers.index("budget_ok")
    assert any(row[model_index] == "Qwen/Qwen2.5-3B-Instruct" for row in rows[1:])
    assert any(row[budget_index] is True for row in rows[1:])
    error_index = headers.index("error_type")
    assert any(row[error_index] for row in rows[1:])
