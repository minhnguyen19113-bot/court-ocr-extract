from __future__ import annotations

import json

from openpyxl import load_workbook

from court_ocr_extract.final_excel_schema import (
    FINAL_EXCEL_COLUMNS,
    FINAL_EXCEL_SHEET_NAME,
)
from court_ocr_extract.ocr_backends.base import OCRPage, OCRResult
from court_ocr_extract.ocr_cache import OCRCacheRecord
from court_ocr_extract.pre_content_ab import run_pre_content_ab_test
from court_ocr_extract.settings import PipelineSettings


def test_compare_workbook_puts_exact_final_sheet_before_debug_sheets(tmp_path) -> None:
    texts = [
        "Bản án số: 01/2099/HS-ST",
        "thụ lý số 999/2099/TLST-HS ngày 09 tháng 4 năm 2099",
        "Thẩm phán - Chủ tọa phiên tòa: Chủ Tọa Synthetic",
        "Đối với bị cáo:",
        "Người Synthetic A, sinh năm 1990",
        "Chỗ ở: Vùng Synthetic A",
        "Bị hại: Người Synthetic B-có mặt",
        "Địa chỉ: Vùng Synthetic B",
        "NỘI DUNG VỤ ÁN",
    ]
    run_pre_content_ab_test(
        [_record(texts)],
        output_dir=tmp_path,
        settings=PipelineSettings(),
        strategies=["rule_anchor_only"],
    )

    workbook = load_workbook(tmp_path / "compare_summary.xlsx", read_only=True)
    assert workbook.sheetnames[0] == FINAL_EXCEL_SHEET_NAME
    final_sheet = workbook[FINAL_EXCEL_SHEET_NAME]
    assert [cell.value for cell in next(final_sheet.iter_rows())] == FINAL_EXCEL_COLUMNS
    assert final_sheet.max_column == 11
    assert {
        "CASES",
        "DEFENDANTS",
        "PARTICIPANTS",
        "TRIAL_PANEL",
        "ANCHOR_BLOCKS",
        "ANCHOR_WARNINGS",
        "LLM_STATUS",
        "RAW_JSON",
    } <= set(workbook.sheetnames[1:])

    html = (tmp_path / "cases" / "synthetic" / "review.html").read_text(
        encoding="utf-8"
    )
    preview_start = html.index("FINAL EXCEL PREVIEW")
    preview_end = html.index("</section>", preview_start)
    preview = html[preview_start:preview_end]
    assert preview_start < html.index("Anchor segmentation")
    assert preview.count("<th>") == 11
    assert "CONFIDENCE" not in preview


def test_rule_then_llm_strategy_also_writes_final_excel_without_real_llm(tmp_path) -> None:
    texts = [
        "Bản án số: 02/2099/HS-ST",
        "thụ lý số 998/2099/TLST-HS ngày 10 tháng 4 năm 2099",
        "Thẩm phán - Chủ tọa phiên tòa: Chủ Tọa Synthetic",
        "Đối với bị cáo:",
        "Người Synthetic C, sinh năm 1991",
        "Bị hại: Người Synthetic D-có mặt",
        "NỘI DUNG VỤ ÁN",
    ]

    def fake_llm(prompt, text):
        request = json.loads(text)
        if request["type"] == "defendant":
            return {
                "full_name": "Người Synthetic C",
                "birth_date_or_year": "1991",
                "current_address": "Vùng Synthetic C",
            }
        return {
            "role": request["role_hint"],
            "full_name": "Người Synthetic D",
            "address": "Vùng Synthetic D",
        }

    run_pre_content_ab_test(
        [_record(texts)],
        output_dir=tmp_path,
        settings=PipelineSettings(),
        strategies=["rule_then_llm_per_block"],
        llm_callable=fake_llm,
    )

    workbook = load_workbook(tmp_path / "compare_summary.xlsx", read_only=True)
    final_sheet = workbook[FINAL_EXCEL_SHEET_NAME]
    assert workbook.sheetnames[0] == FINAL_EXCEL_SHEET_NAME
    assert final_sheet.max_row == 3
    assert [row[4] for row in final_sheet.iter_rows(min_row=2, values_only=True)] == [
        "Bị cáo",
        "Bị hại",
    ]


def _record(texts: list[str]) -> OCRCacheRecord:
    lines = [
        {"line_id": f"p001_l{index:04d}", "page_number": 1, "text": text}
        for index, text in enumerate(texts, start=1)
    ]
    return OCRCacheRecord(
        "synthetic",
        1,
        "synthetic",
        OCRResult(
            backend="synthetic",
            status="success",
            pages_processed=1,
            marker_found=True,
            marker_page=1,
            text="\n".join(texts),
            pages=[OCRPage(page_index=1, lines=lines)],
        ),
    )
