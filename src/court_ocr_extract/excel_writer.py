from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from court_ocr_extract.final_excel_builder import make_final_excel_row
from court_ocr_extract.final_excel_role_policy import (
    classify_final_role,
    normalized_row_identity,
)
from court_ocr_extract.final_excel_schema import (
    FINAL_EXCEL_COLUMNS,
    FINAL_EXCEL_SHEET_NAME,
)
from court_ocr_extract.extractors.rule_parser import parse_vietnamese_date
from court_ocr_extract.models import ExtractionResult
from court_ocr_extract.other_participants_builder import (
    OTHER_PARTICIPANT_COLUMNS,
    OTHER_PARTICIPANTS_SHEET_NAME,
    make_other_participant_row,
)
from court_ocr_extract.source_region_policy import FRONT_PRE_CONTENT
from court_ocr_extract.validation import row_needs_review


EXCEL_HEADERS = FINAL_EXCEL_COLUMNS


def rows_from_payload(case_id: str, payload: dict[str, Any]) -> list[dict[str, Any]]:
    case = payload.get("case") or {}
    participants = payload.get("participants") or []
    rows: list[dict[str, Any]] = []
    seen: set[tuple[str, str, str]] = set()
    document_note = "; ".join(payload.get("document_warnings", []))
    for participant in participants:
        if str(participant.get("source_region") or "") not in {"", FRONT_PRE_CONTENT}:
            continue
        decision = classify_final_role(participant.get("procedural_role"))
        if not decision.include_in_final:
            continue
        identity = normalized_row_identity(
            case_id,
            participant.get("full_name"),
            decision.normalized_role,
        )
        if participant.get("full_name") and identity in seen:
            continue
        seen.add(identity)
        participant_note = "; ".join(participant.get("warnings", []))
        note = "; ".join(item for item in [document_note, participant_note] if item)
        row = _complete_final_row(
            {
                "LOẠI ÁN": case.get("case_type"),
                "SỐ BẢN ÁN": case.get("judgment_number"),
                "NGÀY TUYÊN ÁN (DD/MM/YYYY)": case.get("judgment_date"),
                "SỐ THỤ LÝ": case.get("filing_number"),
                "NGÀY THỤ LÝ (DD/MM/YYYY)": case.get("filing_date"),
                "QUAN HỆ PHÁP LUẬT": case.get("legal_relationship"),
                "TƯ CÁCH TỐ TỤNG": decision.normalized_role,
                "HỌ TÊN ĐƯƠNG SỰ": participant.get("full_name"),
                "NĂM SINH": participant.get("birth_year"),
                "CCCD": participant.get("id_number"),
                "ĐỊA CHỈ": participant.get("address"),
                "HỌ TÊN CHỦ TỌA": case.get("presiding_judge"),
            },
            notes=[note],
        )
        row["_case_id"] = case_id
        row["_needs_review"] = row_needs_review(payload, participant)
        rows.append(row)
    return rows


def rows_from_result(result: ExtractionResult) -> list[dict[str, str]]:
    participants = result.participants
    rows: list[dict[str, str]] = []
    seen: set[tuple[str, str, str]] = set()
    case_id = result.source_file or ""
    common_note = "; ".join(result.warnings)
    for participant in participants:
        decision = classify_final_role(participant.tu_cach_to_tung)
        if not decision.include_in_final:
            continue
        identity = normalized_row_identity(
            case_id,
            participant.ho_ten,
            decision.normalized_role,
        )
        if participant.ho_ten and identity in seen:
            continue
        seen.add(identity)
        note = "; ".join(item for item in [participant.ghi_chu, common_note] if item)
        rows.append(
            _complete_final_row(
                {
                    "LOẠI ÁN": result.case_info.loai_an,
                    "SỐ BẢN ÁN": getattr(result.case_info, "so_ban_an", None),
                    "NGÀY TUYÊN ÁN (DD/MM/YYYY)": getattr(
                        result.case_info,
                        "ngay_tuyen_an",
                        None,
                    ),
                    "SỐ THỤ LÝ": result.case_info.so_thu_ly,
                    "NGÀY THỤ LÝ (DD/MM/YYYY)": result.case_info.ngay_thu_ly,
                    "QUAN HỆ PHÁP LUẬT": result.case_info.quan_he_phap_luat,
                    "TƯ CÁCH TỐ TỤNG": decision.normalized_role,
                    "HỌ TÊN ĐƯƠNG SỰ": participant.ho_ten,
                    "NĂM SINH": participant.nam_sinh,
                    "CCCD": participant.cccd,
                    "ĐỊA CHỈ": participant.dia_chi,
                    "HỌ TÊN CHỦ TỌA": result.case_info.chu_toa,
                },
                notes=[note],
            )
        )
    return rows


def other_rows_from_payload(case_id: str, payload: dict[str, Any]) -> list[dict[str, str]]:
    case = payload.get("case") or {}
    rows: list[dict[str, str]] = []
    seen: set[tuple[str, str, str]] = set()
    for index, participant in enumerate(payload.get("participants") or []):
        if str(participant.get("source_region") or "") not in {"", FRONT_PRE_CONTENT}:
            continue
        decision = classify_final_role(participant.get("procedural_role"))
        if not decision.include_in_other:
            continue
        full_name = participant.get("full_name")
        identity = normalized_row_identity(case_id, full_name, decision.normalized_role)
        if full_name and identity in seen:
            continue
        seen.add(identity if full_name else (*identity[:2], f"{identity[2]}:{index}"))
        notes = [
            "; ".join(participant.get("warnings", [])),
            participant.get("relationship_or_note") or participant.get("relationship"),
        ]
        if decision.category == "other_unclassified":
            notes.append("Tư cách cần review")
        rows.append(
            make_other_participant_row(
                {
                    "LOẠI ÁN": case.get("case_type"),
                    "SỐ THỤ LÝ": case.get("filing_number"),
                    "TƯ CÁCH TỐ TỤNG": participant.get("procedural_role"),
                    "HỌ TÊN": full_name,
                    "NGƯỜI ĐƯỢC ĐẠI DIỆN/BẢO VỆ": participant.get(
                        "represented_person"
                    ),
                    "ĐỊA CHỈ": participant.get("address"),
                    "TÌNH TRẠNG THAM GIA": participant.get("presence_status"),
                },
                notes=notes,
            )
        )
    return rows


def other_rows_from_result(result: ExtractionResult) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    seen: set[tuple[str, str, str]] = set()
    case_id = result.source_file or ""
    for index, participant in enumerate(result.participants):
        decision = classify_final_role(participant.tu_cach_to_tung)
        if not decision.include_in_other:
            continue
        identity = normalized_row_identity(
            case_id,
            participant.ho_ten,
            decision.normalized_role,
        )
        if participant.ho_ten and identity in seen:
            continue
        seen.add(identity if participant.ho_ten else (*identity[:2], f"{identity[2]}:{index}"))
        notes = [participant.ghi_chu]
        if decision.category == "other_unclassified":
            notes.append("Tư cách cần review")
        rows.append(
            make_other_participant_row(
                {
                    "LOẠI ÁN": result.case_info.loai_an,
                    "SỐ THỤ LÝ": result.case_info.so_thu_ly,
                    "TƯ CÁCH TỐ TỤNG": participant.tu_cach_to_tung,
                    "HỌ TÊN": participant.ho_ten,
                    "ĐỊA CHỈ": participant.dia_chi,
                },
                notes=notes,
            )
        )
    return rows


def write_excel(
    draft_records: list[dict[str, Any]],
    output_path: str | Path,
    *,
    run_summary: dict[str, Any] | None = None,
    include_other_participants_output: bool = False,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = FINAL_EXCEL_SHEET_NAME
    data_sheet.append(EXCEL_HEADERS)
    for draft in draft_records:
        for row in rows_from_payload(draft["case_id"], draft["payload"]):
            data_sheet.append([row.get(header) for header in EXCEL_HEADERS])
    format_final_excel_sheet(data_sheet)

    if include_other_participants_output:
        other_sheet = workbook.create_sheet(OTHER_PARTICIPANTS_SHEET_NAME)
        other_sheet.append(OTHER_PARTICIPANT_COLUMNS)
        for draft in draft_records:
            for row in other_rows_from_payload(draft["case_id"], draft["payload"]):
                other_sheet.append([row[column] for column in OTHER_PARTICIPANT_COLUMNS])
        format_other_participants_sheet(other_sheet)

    summary_sheet = workbook.create_sheet("RUN_SUMMARY")
    summary = run_summary or build_run_summary(draft_records)
    for key, value in summary.items():
        summary_sheet.append([key, value])
    summary_sheet.column_dimensions["A"].width = 34
    summary_sheet.column_dimensions["B"].width = 42

    workbook.save(output_path)
    return output_path


def write_excel_from_results(
    results: list[ExtractionResult],
    output_path: str | Path,
    *,
    include_other_participants_output: bool = False,
) -> Path:
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook = Workbook()
    data_sheet = workbook.active
    data_sheet.title = FINAL_EXCEL_SHEET_NAME
    data_sheet.append(EXCEL_HEADERS)
    for result in results:
        for row in rows_from_result(result):
            data_sheet.append([row.get(header) for header in EXCEL_HEADERS])
    format_final_excel_sheet(data_sheet)
    if include_other_participants_output:
        other_sheet = workbook.create_sheet(OTHER_PARTICIPANTS_SHEET_NAME)
        other_sheet.append(OTHER_PARTICIPANT_COLUMNS)
        for result in results:
            for row in other_rows_from_result(result):
                other_sheet.append([row[column] for column in OTHER_PARTICIPANT_COLUMNS])
        format_other_participants_sheet(other_sheet)
    workbook.save(output_path)
    return output_path


def build_run_summary(draft_records: list[dict[str, Any]]) -> dict[str, Any]:
    total = len(draft_records)
    failed = sum(1 for record in draft_records if record.get("status") != "success")
    marker_not_found = sum(1 for record in draft_records if not record.get("marker_found"))
    rows = [row for record in draft_records for row in rows_from_payload(record["case_id"], record["payload"])]
    review_rows = sum(1 for row in rows if row.get("_needs_review"))
    return {
        "Tổng số PDF": total,
        "Số file xử lý thành công": total - failed,
        "Số file lỗi": failed,
        "Số dòng cần review": review_rows,
        "Số file không tìm thấy marker": marker_not_found,
        "Số dòng thiếu thông tin quan trọng": _missing_important_count(rows),
        "OCR backend": _joined_unique(record.get("ocr_backend") for record in draft_records),
        "Extractor backend": _joined_unique(record.get("extractor_backend") for record in draft_records),
        "Debug visual run id": _joined_unique(record.get("debug_run_id") for record in draft_records),
    }


def format_final_excel_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True, color="000000")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [16, 20, 24, 20, 22, 32, 28, 28, 12, 18, 48, 28, 44]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def format_other_participants_sheet(sheet) -> None:
    header_fill = PatternFill("solid", fgColor="E2F0D9")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = Font(bold=True, color="000000")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    widths = [16, 20, 38, 28, 36, 44, 24, 44]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions


def _missing_important_count(rows: list[dict[str, Any]]) -> int:
    important = [
        "SỐ BẢN ÁN",
        "NGÀY TUYÊN ÁN (DD/MM/YYYY)",
        "SỐ THỤ LÝ",
        "NGÀY THỤ LÝ (DD/MM/YYYY)",
        "TƯ CÁCH TỐ TỤNG",
        "HỌ TÊN ĐƯƠNG SỰ",
    ]
    return sum(1 for row in rows if any(not row.get(key) for key in important))


def _joined_unique(values) -> str:
    clean = [str(value) for value in values if value]
    return ", ".join(sorted(set(clean)))


def _complete_final_row(
    values: dict[str, Any],
    *,
    notes: list[str],
) -> dict[str, str]:
    values = dict(values)
    values["NĂM SINH"] = _year_only(values.get("NĂM SINH"))
    for date_field in (
        "NGÀY TUYÊN ÁN (DD/MM/YYYY)",
        "NGÀY THỤ LÝ (DD/MM/YYYY)",
    ):
        raw_date = str(values.get(date_field) or "")
        normalized_date = parse_vietnamese_date(raw_date)
        values[date_field] = normalized_date or ""
        if raw_date and not normalized_date:
            notes.append("Field bị validator loại")
    identity = re.sub(r"\D", "", str(values.get("CCCD") or ""))
    if not 9 <= len(identity) <= 12:
        if identity:
            notes.append("Field bị validator loại")
        identity = ""
    values["CCCD"] = identity
    values["ĐỊA CHỈ"] = re.sub(
        r"^\s*hiện\s+tại\s*[:：]\s*",
        "",
        str(values.get("ĐỊA CHỈ") or ""),
        flags=re.IGNORECASE,
    ).strip(" ;,.")

    if not values.get("LOẠI ÁN"):
        notes.append("Không xác định chắc loại án")
    if not values.get("SỐ BẢN ÁN"):
        notes.append("Thiếu số bản án")
    if not values.get("NGÀY TUYÊN ÁN (DD/MM/YYYY)"):
        notes.append("Thiếu ngày tuyên án")
    if not values.get("SỐ THỤ LÝ"):
        notes.append("Thiếu số thụ lý")
    if not values.get("NGÀY THỤ LÝ (DD/MM/YYYY)"):
        notes.append("Thiếu ngày thụ lý")
    if not values.get("QUAN HỆ PHÁP LUẬT"):
        if "hình sự" in str(values.get("LOẠI ÁN") or "").casefold():
            notes.append("Chưa trích xuất tội danh")
        else:
            notes.append("Không xác định chắc quan hệ pháp luật")
    if not values.get("TƯ CÁCH TỐ TỤNG"):
        notes.append("Thiếu tư cách tố tụng")
    if not values.get("HỌ TÊN ĐƯƠNG SỰ"):
        notes.append("Thiếu họ tên")
    if not values.get("NĂM SINH"):
        notes.append("Thiếu năm sinh")
    if not identity:
        notes.append("Thiếu CCCD/CMND")
    if not values.get("ĐỊA CHỈ"):
        notes.append("Thiếu địa chỉ")
    if not values.get("HỌ TÊN CHỦ TỌA"):
        notes.append("Thiếu chủ tọa")
    return make_final_excel_row(values, notes=notes)


def _year_only(value: Any) -> str:
    years = re.findall(r"(?<!\d)((?:18|19|20)\d{2})(?!\d)", str(value or ""))
    return years[-1] if years else ""
