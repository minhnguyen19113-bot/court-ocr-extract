from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from court_ocr_extract.excel_writer import EXCEL_HEADERS
from court_ocr_extract.final_excel_schema import FINAL_EXCEL_SHEET_NAME
from court_ocr_extract.validation import STATUS_PHRASES, VALID_ROLES


def qa_excel(excel_path: str | Path) -> dict[str, Any]:
    workbook = load_workbook(excel_path, read_only=True, data_only=True)
    if FINAL_EXCEL_SHEET_NAME in workbook.sheetnames:
        sheet = workbook[FINAL_EXCEL_SHEET_NAME]
    elif "DATA" in workbook.sheetnames:
        sheet = workbook["DATA"]
    else:
        sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return {"total rows": 0}
    headers = [str(value) if value is not None else "" for value in rows[0]]
    data = [dict(zip(headers, row, strict=False)) for row in rows[1:]]
    summary = {
        "total rows": len(data),
        "participant rows": sum(1 for row in data if row.get("HỌ TÊN ĐƯƠNG SỰ")),
        "rows need review": sum(1 for row in data if row.get("GHI CHÚ")),
        "invalid date count": sum(1 for row in data if row.get("NGÀY THỤ LÝ (DD/MM/YYYY)") and not _valid_date(str(row.get("NGÀY THỤ LÝ (DD/MM/YYYY)")))),
        "invalid id count": sum(
            1
            for row in data
            if row.get("CCCD")
            and not 9 <= len(re.sub(r"\D", "", str(row.get("CCCD")))) <= 12
        ),
        "suspicious name phrase count": sum(1 for row in data if _suspicious_name(row.get("HỌ TÊN ĐƯƠNG SỰ"))),
        "unknown role count": sum(1 for row in data if row.get("TƯ CÁCH TỐ TỤNG") and row.get("TƯ CÁCH TỐ TỤNG") not in VALID_ROLES),
    }
    for header in EXCEL_HEADERS:
        empty = sum(1 for row in data if not row.get(header))
        summary[f"empty rate {header}"] = round(empty / len(data), 4) if data else 0
    return summary


def print_safe_qa(summary: dict[str, Any]) -> None:
    for key, value in summary.items():
        print(f"{key}: {value}")


def _valid_date(value: str) -> bool:
    return bool(re.fullmatch(r"\d{2}/\d{2}/\d{4}", value))


def _suspicious_name(value: Any) -> bool:
    if not value:
        return False
    lowered = str(value).casefold()
    return any(phrase.casefold() in lowered for phrase in STATUS_PHRASES)
