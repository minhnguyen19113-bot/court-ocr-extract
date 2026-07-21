from __future__ import annotations

import json
from copy import deepcopy
from pathlib import Path
from typing import Any

from openpyxl import Workbook

from court_ocr_extract.charge_parser import (
    iter_verdict_blocks,
    parse_explicit_decision_charges,
)
from court_ocr_extract.decision_tail import DecisionTailRecord
from court_ocr_extract.excel_writer import (
    format_final_excel_sheet,
    format_other_participants_sheet,
)
from court_ocr_extract.extractors.hybrid_pre_content_extractor import HybridPreContentExtractor
from court_ocr_extract.extractors.llm_only_pre_content_extractor import LLMCallable, LLMOnlyPreContentExtractor
from court_ocr_extract.extractors.pre_content_anchor_segmenter import segment_pre_content_anchors
from court_ocr_extract.extractors.pre_content_schema import empty_pre_content_output, flatten_output
from court_ocr_extract.extractors.pre_content_segmenter import segment_pre_content
from court_ocr_extract.extractors.rule_anchor_extractor import extract_rule_anchor_output
from court_ocr_extract.extractors.rule_anchor_strategies import (
    RULE_ANCHOR_STRATEGIES,
    run_rule_anchor_strategy,
)
from court_ocr_extract.extractors.rule_based_pre_content_extractor import extract_pre_content_rules
from court_ocr_extract.final_excel_builder import build_final_excel_rows
from court_ocr_extract.final_excel_role_policy import classify_final_role
from court_ocr_extract.final_excel_schema import (
    FINAL_EXCEL_COLUMNS,
    FINAL_EXCEL_SHEET_NAME,
)
from court_ocr_extract.ocr_cache import OCRCacheRecord
from court_ocr_extract.other_participants_builder import (
    OTHER_PARTICIPANT_COLUMNS,
    OTHER_PARTICIPANTS_SHEET_NAME,
    build_other_participant_rows,
)
from court_ocr_extract.settings import PipelineSettings
from court_ocr_extract.sentence_parser import (
    SENTENCE_REVIEW_WARNINGS,
    empty_sentence_output,
    parse_defendant_sentences,
)
from court_ocr_extract.source_region_policy import (
    DECISION_TAIL,
    FRONT_PRE_CONTENT,
    build_extraction_source_region_audit,
)
from court_ocr_extract.visual_debug import escape, write_html


LEGACY_STRATEGIES = ("hybrid_rule_llm", "legacy_hybrid_rule_llm", "llm_only")
STRATEGIES = LEGACY_STRATEGIES + RULE_ANCHOR_STRATEGIES
DEFAULT_STRATEGIES = ("rule_anchor_only", "rule_then_llm_per_block")

CASES_HEADERS = (
    "case_id", "source_file", "document_type", "ocr_status", "marker_found", "marker_page",
    "early_stop_triggered", "pages_processed", "pages_total", "strategy_used",
    "decision_tail_status", "decision_heading_found", "decision_heading_page",
    "decision_tail_line_count", "verdict_candidate_count", "parsed_charge_count",
    "mapped_defendant_count", "unmapped_verdict_count", "invalid_charge_count",
    "case_charges", "charge_warnings", "court_name",
    "judgment_number", "judgment_date", "case_type", "legal_relationship",
    "case_acceptance_number", "case_acceptance_date", "trial_decision_number",
    "postponement_decision_number", "trial_date_or_location_sentence", "presiding_judge", "clerk",
    "prosecutor", "needs_review", "warnings",
)
DEFENDANTS_HEADERS = (
    "case_id", "strategy", "defendant_index", "entity_id", "source_region",
    "full_name", "alias", "birth_date_or_year",
    "birth_place", "cccd", "permanent_address", "current_address", "detention_status", "presence_status",
    "occupation", "education", "nationality", "ethnicity", "religion", "gender", "father_name",
    "mother_name", "spouse", "children", "criminal_record", "evidence_line_ids", "evidence_text",
    "needs_review", "warnings",
)
PARTICIPANTS_HEADERS = (
    "case_id", "strategy", "participant_index", "entity_id", "source_region",
    "role", "full_name", "birth_date_or_year", "cccd",
    "address", "presence_status", "relationship_or_note", "evidence_line_ids", "evidence_text",
    "needs_review", "warnings",
)
TRIAL_PANEL_HEADERS = (
    "case_id", "strategy", "role", "name", "title_or_position", "organization",
    "evidence_line_ids", "evidence_text", "warnings",
)
LLM_STATUS_HEADERS = (
    "case_id", "strategy", "chunk_name", "block_type", "llm_required", "llm_available",
    "llm_actually_called", "provider", "model", "base_url", "context_window", "input_chars",
    "input_tokens_estimated", "max_output_tokens", "budget_ok", "truncated", "chunked",
    "request_ok", "response_ok", "error_type", "error_message", "duration_ms",
)


def run_pre_content_ab_test(
    records: list[OCRCacheRecord],
    *,
    output_dir: str | Path,
    settings: PipelineSettings,
    strategies: list[str] | None = None,
    limit: int | None = None,
    llm_callable: LLMCallable | None = None,
    llm_preflight: dict[str, Any] | None = None,
    llm_available: bool | None = None,
    decision_tail_records: dict[str, DecisionTailRecord] | None = None,
    include_other_participants_output: bool = False,
) -> dict[str, Any]:
    selected = records[:limit] if limit is not None else records
    requested = list(strategies or DEFAULT_STRATEGIES)
    unknown = [name for name in requested if name not in STRATEGIES]
    if unknown:
        raise ValueError(f"Unsupported pre-content strategies: {', '.join(unknown)}")
    if len(requested) < 1:
        raise ValueError("At least one pre-content strategy is required.")

    available = True if llm_available is None else llm_available
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)
    hybrid = HybridPreContentExtractor(settings, llm_callable=llm_callable)
    llm_only = LLMOnlyPreContentExtractor(settings, llm_callable=llm_callable)
    cases: list[dict[str, Any]] = []

    for record in selected:
        case_dir = output_dir / "cases" / record.case_id
        case_dir.mkdir(parents=True, exist_ok=True)
        segment = segment_pre_content(record.result)
        anchor = segment_pre_content_anchors(segment, case_id=record.case_id)
        legacy_rule = extract_pre_content_rules(segment)
        correction = segment["document_type"] == "correction_notice"
        outputs: dict[str, dict[str, Any]] = {}

        for strategy in requested:
            if correction:
                outputs[strategy] = _skipped_output(
                    segment,
                    strategy,
                    "correction_notice_not_in_judgment_benchmark",
                )
            elif strategy == "rule_anchor_only":
                outputs[strategy] = _rule_anchor_output(anchor, record.case_id, strategy)
            elif strategy in RULE_ANCHOR_STRATEGIES:
                if available:
                    outputs[strategy] = run_rule_anchor_strategy(
                        segment,
                        case_id=record.case_id,
                        strategy=strategy,
                        settings=settings,
                        llm_callable=llm_callable,
                    )
                else:
                    outputs[strategy] = _anchor_llm_unavailable(
                        anchor,
                        record.case_id,
                        strategy,
                        llm_preflight,
                    )
            elif strategy in {"hybrid_rule_llm", "legacy_hybrid_rule_llm"}:
                if available:
                    output = hybrid.extract(segment, case_id=record.case_id)
                    output["strategy"] = strategy
                    outputs[strategy] = output
                else:
                    outputs[strategy] = _hybrid_without_llm(
                        legacy_rule,
                        llm_preflight,
                        strategy=strategy,
                    )
            elif strategy == "llm_only":
                outputs[strategy] = (
                    llm_only.extract(segment, case_id=record.case_id)
                    if available
                    else _llm_unavailable_output(segment, llm_preflight, strategy=strategy)
                )
                outputs[strategy]["strategy"] = strategy

        if decision_tail_records is not None:
            tail_record = decision_tail_records.get(record.case_id)
            for output in outputs.values():
                _attach_decision_tail_charges(output, tail_record, settings=settings)
        for output in outputs.values():
            output["source_region_audit"] = build_extraction_source_region_audit(output)

        left_name = requested[0]
        right_name = requested[1] if len(requested) > 1 else requested[0]
        compare = compare_strategy_outputs(
            outputs[left_name],
            outputs[right_name],
            left_strategy=left_name,
            right_strategy=right_name,
        )
        metrics = _case_metrics(segment, outputs, compare, left_name, right_name)
        case = {
            "case_id": record.case_id,
            "segmenter": segment,
            "anchor_segments": anchor,
            "ocr": _ocr_summary(record),
            "rule_output": legacy_rule,
            "strategy_outputs": outputs,
            "compare": compare,
            "metrics": metrics,
            "benchmark_included": not correction and segment["document_type"] == "judgment_criminal_first_instance",
            "decision_tail": _decision_tail_summary(
                decision_tail_records.get(record.case_id)
                if decision_tail_records is not None
                else None
            ),
        }
        _write_case(
            case_dir,
            case,
            include_other_participants_output=include_other_participants_output,
        )
        cases.append(case)

    summary = _summary(cases, requested, llm_preflight)
    _write_json(output_dir / "compare_summary.json", summary)
    _write_workbook(
        output_dir / "compare_summary.xlsx",
        cases,
        summary,
        include_other_participants_output=include_other_participants_output,
    )
    _write_index(
        output_dir / "index.html",
        cases,
        llm_preflight,
        include_other_participants_output=include_other_participants_output,
    )
    return summary


def compare_strategy_outputs(
    left: dict[str, Any],
    right: dict[str, Any],
    *,
    left_strategy: str,
    right_strategy: str,
) -> dict[str, Any]:
    left_flat, right_flat = flatten_output(left), flatten_output(right)
    rows = []
    for field in sorted(set(left_flat) | set(right_flat)):
        left_value, right_value = left_flat.get(field), right_flat.get(field)
        if left_value in (None, "", []) and right_value in (None, "", []):
            status = "missing_both"
        elif left_value == right_value:
            status = "same"
        elif left_value in (None, "", []) or right_value in (None, "", []):
            status = "missing_one"
        else:
            status = "different"
        rows.append({"field": field, "left": left_value, "right": right_value, "status": status})
    return {
        "left_strategy": left_strategy,
        "right_strategy": right_strategy,
        "fields": rows,
        "disagreement_count": sum(row["status"] in {"different", "missing_one"} for row in rows),
        "conflict_count": len(left.get("conflicts", [])),
    }


def compare_outputs(hybrid: dict[str, Any], llm_only: dict[str, Any]) -> dict[str, Any]:
    """Compatibility wrapper for callers of the legacy two-strategy comparison."""
    result = compare_strategy_outputs(
        hybrid,
        llm_only,
        left_strategy="hybrid_rule_llm",
        right_strategy="llm_only",
    )
    for row in result["fields"]:
        row["hybrid"] = row["left"]
        row["llm_only"] = row["right"]
    return result


def _rule_anchor_output(anchor: dict[str, Any], case_id: str, strategy: str) -> dict[str, Any]:
    output = extract_rule_anchor_output(anchor)
    output.update(case_id=case_id, strategy=strategy)
    return output


def _case_metrics(segment, outputs, compare, left_name, right_name) -> dict[str, Any]:
    primary = outputs[left_name]
    primary_flat = flatten_output(primary)
    present = [field for field, value in primary_flat.items() if value not in (None, "", [])]
    evidence_fields = {item.get("field") for item in primary.get("evidence", []) if item.get("field")}
    statuses = [status for output in outputs.values() for status in output.get("llm_status", [])]
    secondary = outputs[right_name]
    return {
        "primary_strategy": left_name,
        "secondary_strategy": right_name,
        "field_present_count": len(present),
        "field_missing_count": sum(value in (None, "", []) for value in primary_flat.values()),
        "conflict_count": len(primary.get("conflicts", [])),
        "needs_review_count": sum(int(bool(output.get("needs_review"))) for output in outputs.values()),
        "evidence_coverage_rate": round(len(evidence_fields & set(present)) / max(1, len(present)), 4),
        "participant_count": len(primary.get("participants", [])),
        "defendant_count": len(primary.get("defendants", [])),
        "document_type_detected": segment.get("document_type"),
        "segment_stop_found": bool(segment.get("stop_line_id")),
        "llm_json_valid": bool(secondary.get("llm_json_valid", False)),
        "secondary_status": secondary.get("status", "unknown"),
        "llm_only_status": outputs.get("llm_only", secondary).get("status", "not_requested"),
        "llm_actually_called": any(item.get("llm_actually_called") for item in statuses),
        "llm_chunk_count": len(statuses),
        "llm_chunk_error_count": sum(bool(item.get("error_type")) for item in statuses),
        "strategy_disagreement_count": compare["disagreement_count"],
        "hybrid_vs_llm_disagreement_count": compare["disagreement_count"],
    }


def _summary(cases, strategies, llm_preflight):
    return {
        "task": "pre_content_rule_anchor_comparison",
        "strategies": strategies,
        "llm_preflight": llm_preflight,
        "case_count": len(cases),
        "judgment_benchmark_count": sum(case["benchmark_included"] for case in cases),
        "correction_notice_count": sum(case["segmenter"]["document_type"] == "correction_notice" for case in cases),
        "unknown_document_count": sum(case["segmenter"]["document_type"] == "unknown" for case in cases),
        "totals": {
            key: sum(case["metrics"].get(key, 0) for case in cases)
            for key in (
                "field_present_count", "field_missing_count", "conflict_count", "needs_review_count",
                "participant_count", "defendant_count", "llm_chunk_count", "llm_chunk_error_count",
                "strategy_disagreement_count",
            )
        },
        "cases": [
            {"case_id": case["case_id"], "benchmark_included": case["benchmark_included"], **case["metrics"]}
            for case in cases
        ],
    }


def _write_case(
    case_dir: Path,
    case: dict[str, Any],
    *,
    include_other_participants_output: bool,
) -> None:
    (case_dir / "pre_content_text.md").write_text(case["segmenter"]["pre_content_text"], encoding="utf-8")
    for key, filename in (
        ("segmenter", "segmenter.json"),
        ("anchor_segments", "anchor_segments.json"),
        ("rule_output", "rule_output.json"),
        ("compare", "compare.json"),
    ):
        _write_json(case_dir / filename, case[key])
    for strategy, output in case["strategy_outputs"].items():
        _write_json(case_dir / f"{strategy}_output.json", output)
        if strategy in {"hybrid_rule_llm", "legacy_hybrid_rule_llm"}:
            _write_json(case_dir / "hybrid_output.json", output)
        elif strategy == "llm_only":
            _write_json(case_dir / "llm_only_output.json", output)
    _write_case_review(
        case_dir / "review.html",
        case,
        include_other_participants_output=include_other_participants_output,
    )


def _write_case_review(
    path: Path,
    case: dict[str, Any],
    *,
    include_other_participants_output: bool,
) -> None:
    compare = case["compare"]
    compare_rows = "".join(
        f'<tr><td>{escape(item["field"])}</td><td>{escape(_display(item["left"]))}</td>'
        f'<td>{escape(_display(item["right"]))}</td><td>{escape(item["status"])}</td></tr>'
        for item in compare["fields"]
    )
    source_lines = "".join(
        f'<div id="{escape(line.get("line_id"))}"><code>{escape(line.get("line_id"))}</code> '
        f'{escape(line.get("text"))}</div>'
        for line in case["segmenter"].get("pre_content_lines", [])
    )
    output_panels = "".join(
        f'<section><h2>{escape(strategy)}</h2><pre>{escape(_display(output))}</pre>'
        f'{_evidence_table(output, strategy)}</section>'
        for strategy, output in case["strategy_outputs"].items()
    )
    body = (
        f'<h1>{escape(case["case_id"])}</h1>'
        + _final_excel_preview_html(_final_rows_for_case(case))
        + (
            _other_participants_preview_html(_other_rows_for_case(case))
            if include_other_participants_output
            else ""
        )
        + _charge_summary_html(case)
        + _defendant_charge_map_html(case)
        + _source_region_audit_html(case)
        + _llm_runtime_html(case)
        + '<section><h2>Văn bản pre-content</h2>' + source_lines + '</section>'
        + _anchor_html(case["anchor_segments"])
        + _metadata_anchor_evidence_html(case)
        + _role_policy_html(case)
        + _charge_evidence_html(case)
        + output_panels
        + '<h2>So sánh field</h2><table><tr><th>Field</th>'
        + f'<th>{escape(compare["left_strategy"])}</th><th>{escape(compare["right_strategy"])}</th><th>Status</th></tr>'
        + compare_rows + '</table>'
    )
    write_html(path, "Rule anchor pre-content review", body)


def _anchor_html(anchor: dict[str, Any]) -> str:
    blocks = []
    for kind, values in (
        ("Bị cáo", anchor.get("defendant_blocks", [])),
        ("Người tham gia tố tụng", anchor.get("participant_blocks", [])),
    ):
        for block in values:
            line_ids = ", ".join(block.get("line_ids", []))
            blocks.append(
                f'<h3>{escape(kind)}: {escape(block.get("block_id"))}</h3>'
                f'<p>Role: {escape(block.get("role_hint"))}; line ids: {escape(line_ids)}; '
                f'split: {escape(block.get("split_reason"))}</p>'
                f'<pre>{escape(block.get("text"))}</pre>'
            )
    return (
        '<section><h2>Anchor segmentation</h2>'
        f'<h3>Metadata lines</h3><pre>{escape(_display(anchor.get("metadata_lines", [])))}</pre>'
        f'<h3>Trial panel lines</h3><pre>{escape(_display(anchor.get("trial_panel_lines", [])))}</pre>'
        f'<h3>Defendant region start/end</h3><pre>{escape(_display(anchor.get("defendant_region", {})))}</pre>'
        f'<h3>Rejected defendant candidates</h3><pre>{escape(_display(anchor.get("rejected_defendant_candidates", [])))}</pre>'
        + "".join(blocks)
        + f'<h3>Validator/cảnh báo</h3><pre>{escape(_display(anchor.get("warnings", [])))}</pre></section>'
    )


def _metadata_anchor_evidence_html(case: dict[str, Any]) -> str:
    rows = []
    for strategy, output in case.get("strategy_outputs", {}).items():
        for item in output.get("evidence", []):
            if not str(item.get("field") or "").startswith("metadata."):
                continue
            rows.append(
                "<tr>"
                f"<td>{escape(strategy)}</td>"
                f"<td>{escape(item.get('field'))}</td>"
                f"<td>{escape(item.get('value'))}</td>"
                f"<td>{escape(item.get('line_id'))}</td>"
                f"<td>{escape(item.get('text'))}</td>"
                "</tr>"
            )
    return (
        '<section><h2>Metadata anchor evidence</h2>'
        '<table><tr><th>Strategy</th><th>Field</th><th>Value</th>'
        '<th>Line ID</th><th>Evidence</th></tr>'
        + "".join(rows)
        + "</table></section>"
    )


def _llm_runtime_html(case: dict[str, Any]) -> str:
    statuses = [
        status
        for output in case["strategy_outputs"].values()
        for status in output.get("llm_status", [])
    ]
    rows = "".join(
        f'<tr><td>{escape(item.get("strategy"))}</td><td>{escape(item.get("chunk_name"))}</td>'
        f'<td>{escape(item.get("block_type"))}</td><td>{escape(item.get("model"))}</td>'
        f'<td>{escape(item.get("base_url"))}</td><td>{escape(item.get("context_window"))}</td>'
        f'<td>{escape(item.get("estimated_input_tokens"))}</td>'
        f'<td>{escape(item.get("llm_actually_called"))}</td><td>{escape(item.get("error_type"))}</td></tr>'
        for item in statuses
    )
    return (
        '<section><h2>Trạng thái Local LLM</h2><table><tr><th>Strategy</th><th>Block</th>'
        '<th>Loại block</th><th>Model</th><th>Base URL</th><th>Context</th>'
        '<th>Input token ước lượng</th><th>Đã gọi</th><th>Lỗi</th></tr>' + rows + '</table></section>'
    )


def _write_index(
    path: Path,
    cases: list[dict[str, Any]],
    preflight: dict[str, Any] | None,
    *,
    include_other_participants_output: bool,
) -> None:
    runtime = preflight or {"ok": "not_run", "model": "", "base_url": ""}
    rows = "".join(
        f'<tr><td><a href="cases/{escape(case["case_id"])}/review.html">{escape(case["case_id"])}</a></td>'
        f'<td>{escape(case["metrics"]["document_type_detected"])}</td>'
        f'<td>{escape(case["metrics"]["primary_strategy"])}</td>'
        f'<td>{escape(case["metrics"]["secondary_status"])}</td>'
        f'<td>{escape(case["metrics"]["llm_actually_called"])}</td>'
        f'<td>{escape(case["metrics"]["llm_chunk_count"])}</td>'
        f'<td>{escape(case["metrics"]["llm_chunk_error_count"])}</td></tr>'
        for case in cases
    )
    body = (
        '<h1>So sánh extraction pre-content theo anchor</h1>'
        + _final_excel_preview_html(
            [row for case in cases for row in _final_rows_for_case(case)]
        )
        + (
            _other_participants_preview_html(
                [row for case in cases for row in _other_rows_for_case(case)]
            )
            if include_other_participants_output
            else ""
        )
        + "".join(_charge_summary_html(case) for case in cases)
        + "".join(_defendant_charge_map_html(case) for case in cases)
        + "".join(_source_region_audit_html(case) for case in cases)
        + f'<p>LLM preflight: <strong>{escape(runtime.get("ok"))}</strong>; model: {escape(runtime.get("model"))}; '
        + f'base URL: {escape(runtime.get("base_url"))}</p>'
        + '<table><tr><th>Case</th><th>Document type</th><th>Strategy chính</th>'
        + '<th>Trạng thái strategy thứ hai</th><th>Đã gọi LLM</th><th>Số block gọi</th><th>Block lỗi</th></tr>'
        + rows + '</table>'
    )
    write_html(path, "Rule anchor pre-content comparison", body)


def _write_workbook(
    path: Path,
    cases: list[dict[str, Any]],
    summary: dict[str, Any],
    *,
    include_other_participants_output: bool,
) -> None:
    workbook = Workbook()
    final_sheet = workbook.active
    final_sheet.title = FINAL_EXCEL_SHEET_NAME
    final_sheet.append(FINAL_EXCEL_COLUMNS)
    for case in cases:
        for row in _final_rows_for_case(case):
            final_sheet.append([row[column] for column in FINAL_EXCEL_COLUMNS])
    format_final_excel_sheet(final_sheet)

    if include_other_participants_output:
        other_sheet = workbook.create_sheet(OTHER_PARTICIPANTS_SHEET_NAME)
        other_sheet.append(OTHER_PARTICIPANT_COLUMNS)
        for case in cases:
            for row in _other_rows_for_case(case):
                other_sheet.append([row[column] for column in OTHER_PARTICIPANT_COLUMNS])
        format_other_participants_sheet(other_sheet)

    sheet_names = (
        "SUMMARY", "ANCHOR_BLOCKS", "ANCHOR_WARNINGS", "CASES", "DEFENDANTS",
        "PARTICIPANTS", "TRIAL_PANEL", "LLM_STATUS", "FIELD_LONG", "EVIDENCE_LINES",
        "RAW_JSON", "CASE_COMPARE", "CONFLICTS", "MISSING_FIELDS", "NEEDS_REVIEW",
        "DOC_ROUTER", "ROLE_POLICY", "CHARGES", "DEFENDANT_CHARGES",
        "SOURCE_REGION_AUDIT", "CHARGE_WARNINGS", "CHARGE_EVIDENCE",
        "DEFENDANT_SENTENCES", "SENTENCE_EVIDENCE", "SENTENCE_WARNINGS",
    )
    sheets = {name: workbook.create_sheet(name) for name in sheet_names}
    sheets["SUMMARY"].append(["METRIC", "VALUE"])
    for key, value in summary.items():
        if key != "cases":
            sheets["SUMMARY"].append([key, _display(value)])
    for case in cases:
        output = _primary_output_for_case(case)
        charge_output = output.get("charge_output", {})
        charge_warnings = (
            charge_output.get("warnings", [])
            if isinstance(charge_output, dict)
            else []
        )
        sheets["SUMMARY"].append([
            f"charge_warnings:{case['case_id']}",
            "; ".join(str(value) for value in charge_warnings),
        ])
    for name, headers in (
        ("CASES", CASES_HEADERS),
        ("DEFENDANTS", DEFENDANTS_HEADERS),
        ("PARTICIPANTS", PARTICIPANTS_HEADERS),
        ("TRIAL_PANEL", TRIAL_PANEL_HEADERS),
        ("LLM_STATUS", LLM_STATUS_HEADERS),
    ):
        sheets[name].append(list(headers))
    sheets["ANCHOR_BLOCKS"].append([
        "case_id", "block_type", "block_id", "role_hint", "line_ids", "start_line_id",
        "end_line_id", "split_reason", "text",
    ])
    sheets["ANCHOR_WARNINGS"].append(["case_id", "warning"])
    sheets["FIELD_LONG"].append(["case_id", "strategy", "field", "value"])
    sheets["EVIDENCE_LINES"].append(["case_id", "strategy", "field", "value", "line_id", "text"])
    sheets["RAW_JSON"].append(["case_id", "strategy", "json"])
    sheets["CASE_COMPARE"].append(["case_id", "field", "left_strategy", "left", "status", "right_strategy", "right"])
    sheets["CONFLICTS"].append(["case_id", "strategy", "field", "kept_value", "detail"])
    sheets["MISSING_FIELDS"].append(["case_id", "field", "left", "right"])
    sheets["NEEDS_REVIEW"].append(["case_id", "strategy", "warnings"])
    sheets["DOC_ROUTER"].append(["case_id", "document_type", "benchmark_included", "warnings"])
    sheets["ROLE_POLICY"].append([
        "case_id", "strategy", "source_group", "original_role", "normalized_role",
        "category", "include_in_final", "include_in_other", "reason",
    ])
    sheets["CHARGE_EVIDENCE"].append([
        "case_id", "strategy", "charge", "defendant_entity_ids",
        "defendant_names", "source_region", "page_number", "line_ids",
        "match_method", "confidence", "raw_text",
    ])
    sheets["CHARGES"].append([
        "case_id", "charge", "source_region", "page_number", "line_ids",
        "match_method", "confidence", "raw_text",
    ])
    sheets["DEFENDANT_CHARGES"].append([
        "case_id", "defendant_entity_id", "defendant_name", "charge",
        "mapping_status", "mapping_method", "evidence_line_ids",
    ])
    sheets["SOURCE_REGION_AUDIT"].append([
        "case_id", "field_name", "source_region", "source_page",
        "evidence_line_ids", "allowed", "warning",
    ])
    sheets["CHARGE_WARNINGS"].append([
        "case_id", "strategy", "warning",
    ])
    sheets["DEFENDANT_SENTENCES"].append([
        "case_id", "defendant_entity_id", "defendant_name",
        "primary_penalty_kind", "primary_penalty_text", "display_text",
        "suspended", "probation_text", "probation_start_text", "execution_status",
        "sentence_start_text", "detention_credit_text", "completion_text",
        "release_text", "aggregate_penalty_text", "additional_penalties",
        "source_region", "page_number", "line_ids", "confidence", "warnings",
    ])
    sheets["SENTENCE_EVIDENCE"].append([
        "case_id", "evidence_type", "defendant_entity_ids", "defendant_names",
        "raw_text", "page_number", "line_ids", "source_region", "match_method",
        "confidence", "warnings",
    ])
    sheets["SENTENCE_WARNINGS"].append([
        "case_id", "warning", "severity", "scope", "defendant_entity_id",
        "defendant_name", "page_number", "line_ids", "raw_text", "source_region",
        "front_name", "decision_name", "match_method", "similarity",
    ])

    for case in cases:
        _append_anchor_rows(sheets, case)
        _append_structured_rows(sheets, case)
        case_id = case["case_id"]
        sheets["DOC_ROUTER"].append([
            case_id,
            case["segmenter"]["document_type"],
            case["benchmark_included"],
            "; ".join(case["segmenter"]["warnings"]),
        ])
        compare = case["compare"]
        for row in compare["fields"]:
            sheets["CASE_COMPARE"].append([
                case_id, row["field"], compare["left_strategy"], _display(row["left"]),
                row["status"], compare["right_strategy"], _display(row["right"]),
            ])
            if row["status"] in {"missing_both", "missing_one"}:
                sheets["MISSING_FIELDS"].append([
                    case_id, row["field"], _display(row["left"]), _display(row["right"]),
                ])
        for strategy, output in case["strategy_outputs"].items():
            for conflict in output.get("conflicts", []):
                sheets["CONFLICTS"].append([
                    case_id, strategy, conflict.get("field"), _display(conflict.get("kept_value")),
                    _display(conflict),
                ])
    workbook.save(path)


def _final_rows_for_case(case: dict[str, Any]) -> list[dict[str, str]]:
    if case.get("segmenter", {}).get("document_type") == "correction_notice":
        return []
    return build_final_excel_rows(_primary_output_for_case(case))


def _other_rows_for_case(case: dict[str, Any]) -> list[dict[str, str]]:
    if case.get("segmenter", {}).get("document_type") == "correction_notice":
        return []
    return build_other_participant_rows(_primary_output_for_case(case))


def _primary_output_for_case(case: dict[str, Any]) -> dict[str, Any]:
    strategy = str(case.get("metrics", {}).get("primary_strategy") or "")
    outputs = case.get("strategy_outputs", {})
    output = outputs.get(strategy) if isinstance(outputs, dict) else None
    if not isinstance(output, dict):
        output = next(
            (
                value
                for value in outputs.values()
                if isinstance(value, dict)
            ),
            {},
        ) if isinstance(outputs, dict) else {}
    return output


def _final_excel_preview_html(rows: list[dict[str, str]]) -> str:
    headers = "".join(f"<th>{escape(column)}</th>" for column in FINAL_EXCEL_COLUMNS)
    body_rows = "".join(
        "<tr>"
        + "".join(f"<td>{escape(row.get(column))}</td>" for column in FINAL_EXCEL_COLUMNS)
        + "</tr>"
        for row in rows
    )
    if not body_rows:
        body_rows = (
            f'<tr><td colspan="{len(FINAL_EXCEL_COLUMNS)}">Không có dòng final cho tài liệu này.</td></tr>'
        )
    return (
        '<section><h2>FINAL EXCEL PREVIEW</h2>'
        f'<table><tr>{headers}</tr>{body_rows}</table></section>'
    )


def _other_participants_preview_html(rows: list[dict[str, str]]) -> str:
    headers = "".join(
        f"<th>{escape(column)}</th>" for column in OTHER_PARTICIPANT_COLUMNS
    )
    body_rows = "".join(
        "<tr>"
        + "".join(
            f"<td>{escape(row.get(column))}</td>"
            for column in OTHER_PARTICIPANT_COLUMNS
        )
        + "</tr>"
        for row in rows
    )
    if not body_rows:
        body_rows = (
            f'<tr><td colspan="{len(OTHER_PARTICIPANT_COLUMNS)}">'
            "Không có người tham gia khác.</td></tr>"
        )
    return (
        '<section><h2>NGƯỜI THAM GIA KHÁC</h2>'
        f'<table><tr>{headers}</tr>{body_rows}</table></section>'
    )


def _charge_summary_html(case: dict[str, Any]) -> str:
    output = _primary_output_for_case(case)
    charge_output = output.get("charge_output", {})
    charges = charge_output.get("case_charges", []) if isinstance(charge_output, dict) else []
    rows = "".join(
        f"<tr><td>{index}</td><td>{escape(charge)}</td><td>decision_tail</td></tr>"
        for index, charge in enumerate(charges, start=1)
    )
    if not rows:
        rows = '<tr><td colspan="3">Chưa có tội danh explicit từ phần Quyết định.</td></tr>'
    warnings = charge_output.get("warnings", []) if isinstance(charge_output, dict) else []
    candidates = (
        charge_output.get("verdict_candidate_blocks", [])
        if isinstance(charge_output, dict)
        else []
    )
    candidate_rows = "".join(
        "<tr>"
        f"<td>{index}</td>"
        f"<td>{escape(item.get('page_number'))}</td>"
        f"<td>{escape(', '.join(item.get('line_ids', [])))}</td>"
        f"<td><pre>{escape(item.get('raw_text'))}</pre></td>"
        "</tr>"
        for index, item in enumerate(candidates, start=1)
        if isinstance(item, dict)
    )
    if not candidate_rows:
        candidate_rows = '<tr><td colspan="4">Không có verdict candidate.</td></tr>'
    evidence_rows = "".join(
        "<tr>"
        f"<td>{escape(item.get('raw_charge'))}</td>"
        f"<td>{escape(item.get('normalized_charge') or item.get('charge'))}</td>"
        f"<td>{escape(', '.join(item.get('defendant_names', [])))}</td>"
        f"<td>{escape(', '.join(item.get('line_ids', [])))}</td>"
        f"<td>{escape(item.get('raw_text'))}</td>"
        "</tr>"
        for item in charge_output.get("charge_evidence", [])
        if isinstance(charge_output, dict) and isinstance(item, dict)
    )
    if not evidence_rows:
        evidence_rows = '<tr><td colspan="5">Chưa có charge evidence.</td></tr>'
    return (
        '<section><h2>CHARGE SUMMARY</h2>'
        f'<p>Case: {escape(case.get("case_id"))}; decision-tail status: '
        f'{escape(output.get("decision_tail_status") or "not_attached")}; '
        f'heading found: {escape(output.get("decision_heading_found"))}; '
        f'heading page: {escape(output.get("decision_heading_page"))}; '
        f'tail lines: {escape(output.get("decision_tail_line_count"))}; '
        f'verdict candidates: {escape(output.get("verdict_candidate_count"))}; '
        f'parsed charges: {escape(output.get("parsed_charge_count"))}; '
        f'mapped defendants: {escape(output.get("mapped_defendant_count"))}; '
        f'unmapped verdicts: {escape(output.get("unmapped_verdict_count"))}; '
        f'invalid charges: {escape(output.get("invalid_charge_count"))}</p>'
        '<table><tr><th>STT</th><th>Tội danh</th><th>Source region</th></tr>'
        + rows
        + '</table><h3>Verdict candidate blocks</h3>'
        + '<table><tr><th>STT</th><th>Page</th><th>Line IDs</th><th>Raw block</th></tr>'
        + candidate_rows
        + '</table><h3>Charge evidence</h3>'
        + '<table><tr><th>Raw charge</th><th>Normalized charge</th>'
        '<th>Bị cáo</th><th>Line IDs</th><th>Evidence</th></tr>'
        + evidence_rows
        + '</table><p>Cảnh báo: '
        + escape("; ".join(str(value) for value in warnings) or "Không")
        + "</p></section>"
    )


def _defendant_charge_map_html(case: dict[str, Any]) -> str:
    output = _primary_output_for_case(case)
    charge_output = output.get("charge_output", {})
    mapping = charge_output.get("defendant_charge_map", {}) if isinstance(charge_output, dict) else {}
    evidence = charge_output.get("charge_evidence", []) if isinstance(charge_output, dict) else []
    rows = []
    for index, defendant in enumerate(output.get("defendants", []), start=1):
        entity_id = str(
            defendant.get("entity_id")
            or defendant.get("source_block_id")
            or f"defendant_{index:03d}"
        )
        charges = mapping.get(entity_id, []) if isinstance(mapping, dict) else []
        methods = list(
            dict.fromkeys(
                str(item.get("match_method") or "")
                for item in evidence
                if isinstance(item, dict)
                and entity_id in item.get("defendant_entity_ids", [])
                and item.get("match_method")
            )
        )
        rows.append(
            "<tr>"
            f"<td>{escape(entity_id)}</td>"
            f"<td>{escape(defendant.get('full_name'))}</td>"
            f"<td>{escape('; '.join(charges))}</td>"
            f"<td>{'mapped' if charges else 'unmapped'}</td>"
            f"<td>{escape('; '.join(methods))}</td>"
            "</tr>"
        )
    if not rows:
        rows.append('<tr><td colspan="5">Không có defendant front entity.</td></tr>')
    return (
        '<section><h2>DEFENDANT → CHARGE MAP</h2>'
        '<table><tr><th>Entity ID</th><th>Họ tên</th><th>Tội danh</th>'
        '<th>Trạng thái</th><th>Phương pháp</th></tr>'
        + "".join(rows)
        + "</table></section>"
    )


def _source_region_audit_html(case: dict[str, Any]) -> str:
    output = _primary_output_for_case(case)
    rows = "".join(
        "<tr>"
        f"<td>{escape(item.get('field_name'))}</td>"
        f"<td>{escape(item.get('source_region'))}</td>"
        f"<td>{escape(item.get('source_page'))}</td>"
        f"<td>{escape(', '.join(item.get('evidence_line_ids', [])))}</td>"
        f"<td>{escape(item.get('allowed'))}</td>"
        f"<td>{escape(item.get('warning'))}</td>"
        "</tr>"
        for item in output.get("source_region_audit", [])
        if isinstance(item, dict)
    )
    if not rows:
        rows = '<tr><td colspan="6">Chưa có source-region audit row.</td></tr>'
    return (
        '<section><h2>SOURCE REGION AUDIT</h2>'
        '<table><tr><th>Field</th><th>Source region</th><th>Page</th>'
        '<th>Line IDs</th><th>Allowed</th><th>Warning</th></tr>'
        + rows
        + "</table></section>"
    )


def _role_policy_html(case: dict[str, Any]) -> str:
    rows = []
    for strategy, output in case.get("strategy_outputs", {}).items():
        for _defendant in output.get("defendants", []):
            decision = classify_final_role("Bị cáo")
            rows.append(
                "<tr>"
                f"<td>{escape(strategy)}</td>"
                "<td>Bị cáo</td>"
                f"<td>{escape(decision.normalized_role)}</td>"
                f"<td>{escape(decision.category)}</td>"
                f"<td>{escape(decision.include_in_final)}</td>"
                f"<td>{escape(decision.include_in_other)}</td>"
                f"<td>{escape(decision.reason)}</td>"
                "</tr>"
            )
        for participant in output.get("participants", []):
            decision = classify_final_role(participant.get("role"))
            rows.append(
                "<tr>"
                f"<td>{escape(strategy)}</td>"
                f"<td>{escape(participant.get('role'))}</td>"
                f"<td>{escape(decision.normalized_role)}</td>"
                f"<td>{escape(decision.category)}</td>"
                f"<td>{escape(decision.include_in_final)}</td>"
                f"<td>{escape(decision.include_in_other)}</td>"
                f"<td>{escape(decision.reason)}</td>"
                "</tr>"
            )
    return (
        '<section><h2>Quyết định role inclusion/exclusion</h2>'
        '<table><tr><th>Strategy</th><th>Role gốc</th><th>Role chuẩn hóa</th>'
        '<th>Nhóm</th><th>FINAL_EXCEL</th><th>NGUOI_THAM_GIA_KHAC</th>'
        '<th>Lý do</th></tr>' + "".join(rows) + "</table></section>"
    )


def _charge_evidence_html(case: dict[str, Any]) -> str:
    rows = []
    for strategy, output in case.get("strategy_outputs", {}).items():
        charge_output = output.get("charge_output", {})
        for item in charge_output.get("charge_evidence", []):
            rows.append(
                "<tr>"
                f"<td>{escape(strategy)}</td>"
                f"<td>{escape(item.get('raw_charge'))}</td>"
                f"<td>{escape(item.get('normalized_charge') or item.get('charge'))}</td>"
                f"<td>{escape(', '.join(item.get('defendant_names', [])))}</td>"
                f"<td>{escape(', '.join(item.get('line_ids', [])))}</td>"
                f"<td>{escape(item.get('source_region'))}</td>"
                f"<td>{escape(item.get('match_method'))}</td>"
                f"<td>{escape(item.get('raw_text'))}</td>"
                "</tr>"
            )
    return (
        '<section><h2>Charge evidence</h2>'
        '<table><tr><th>Strategy</th><th>Raw charge</th>'
        '<th>Normalized charge</th><th>Bị cáo</th>'
        '<th>Line IDs</th><th>Source region</th><th>Match method</th><th>Evidence</th></tr>'
        + "".join(rows)
        + "</table></section>"
    )


def _append_anchor_rows(sheets, case) -> None:
    case_id = case["case_id"]
    anchor = case["anchor_segments"]
    for block_type, blocks in (
        ("defendant", anchor.get("defendant_blocks", [])),
        ("participant", anchor.get("participant_blocks", [])),
    ):
        for block in blocks:
            sheets["ANCHOR_BLOCKS"].append([
                case_id, block_type, block.get("block_id"), block.get("role_hint"),
                "; ".join(block.get("line_ids", [])), block.get("start_line_id"),
                block.get("end_line_id"), block.get("split_reason"), block.get("text"),
            ])
    for warning in anchor.get("warnings", []):
        sheets["ANCHOR_WARNINGS"].append([case_id, warning])


def _append_structured_rows(sheets, case) -> None:
    case_id = case["case_id"]
    line_text = {
        str(line.get("line_id")): str(line.get("text") or "")
        for line in case["segmenter"].get("pre_content_lines", [])
    }
    for strategy, output in case["strategy_outputs"].items():
        metadata = output.get("metadata", {})
        panel = output.get("trial_panel", {})
        ocr = case["ocr"]
        cases_row = {
            "case_id": case_id,
            "source_file": ocr["source_file"],
            "document_type": output.get("document_type"),
            "ocr_status": ocr["ocr_status"],
            "marker_found": ocr["marker_found"],
            "marker_page": ocr["marker_page"],
            "early_stop_triggered": ocr["early_stop_triggered"],
            "pages_processed": ocr["pages_processed"],
            "pages_total": ocr["pages_total"],
            "strategy_used": strategy,
            "decision_tail_status": output.get("decision_tail_status"),
            "decision_heading_found": output.get("decision_heading_found"),
            "decision_heading_page": output.get("decision_heading_page"),
            "decision_tail_line_count": output.get("decision_tail_line_count"),
            "verdict_candidate_count": output.get("verdict_candidate_count"),
            "parsed_charge_count": output.get("parsed_charge_count"),
            "mapped_defendant_count": output.get("mapped_defendant_count"),
            "unmapped_verdict_count": output.get("unmapped_verdict_count"),
            "invalid_charge_count": output.get("invalid_charge_count"),
            "case_charges": "; ".join(output.get("case_charges", [])),
            "charge_warnings": "; ".join(
                output.get("charge_output", {}).get("warnings", [])
            ),
            **metadata,
            "trial_date_or_location_sentence": metadata.get("trial_location_or_date_sentence"),
            "presiding_judge": panel.get("presiding_judge"),
            "clerk": panel.get("clerk"),
            "prosecutor": panel.get("prosecutor"),
            "needs_review": output.get("needs_review"),
            "warnings": "; ".join(output.get("warnings", [])),
        }
        sheets["CASES"].append([cases_row.get(header) for header in CASES_HEADERS])
        for index, item in enumerate(output.get("defendants", []), start=1):
            row = {"case_id": case_id, "strategy": strategy, "defendant_index": index, **item}
            row["evidence_line_ids"] = "; ".join(item.get("evidence_line_ids", []))
            row["evidence_text"] = _evidence_text(item, line_text)
            row["warnings"] = "; ".join(item.get("warnings", []))
            sheets["DEFENDANTS"].append([row.get(header) for header in DEFENDANTS_HEADERS])
            decision = classify_final_role("Bị cáo")
            sheets["ROLE_POLICY"].append([
                case_id, strategy, "defendant", "Bị cáo",
                decision.normalized_role, decision.category,
                decision.include_in_final, decision.include_in_other, decision.reason,
            ])
        for index, item in enumerate(output.get("participants", []), start=1):
            row = {"case_id": case_id, "strategy": strategy, "participant_index": index, **item}
            row["relationship_or_note"] = item.get("relationship_or_note") or item.get("relationship")
            row["evidence_line_ids"] = "; ".join(item.get("evidence_line_ids", []))
            row["evidence_text"] = _evidence_text(item, line_text)
            row["warnings"] = "; ".join(item.get("warnings", []))
            sheets["PARTICIPANTS"].append([row.get(header) for header in PARTICIPANTS_HEADERS])
            decision = classify_final_role(item.get("role"))
            sheets["ROLE_POLICY"].append([
                case_id, strategy, "participant", item.get("role"),
                decision.normalized_role, decision.category,
                decision.include_in_final, decision.include_in_other, decision.reason,
            ])
        _append_trial_panel(sheets["TRIAL_PANEL"], case_id, strategy, panel)
        for court_role in ("Thẩm phán", "Hội thẩm", "Thư ký", "Kiểm sát viên"):
            decision = classify_final_role(court_role)
            sheets["ROLE_POLICY"].append([
                case_id, strategy, "trial_panel", court_role,
                decision.normalized_role, decision.category,
                decision.include_in_final, decision.include_in_other, decision.reason,
            ])
        for status in output.get("llm_status", []):
            row = {"case_id": case_id, **status}
            row["input_tokens_estimated"] = status.get("estimated_input_tokens")
            sheets["LLM_STATUS"].append([row.get(header) for header in LLM_STATUS_HEADERS])
        for field, value in flatten_output(output).items():
            sheets["FIELD_LONG"].append([case_id, strategy, field, _display(value)])
        for evidence in output.get("evidence", []):
            sheets["EVIDENCE_LINES"].append([
                case_id, strategy, evidence.get("field"), evidence.get("value"),
                evidence.get("line_id"), evidence.get("text"),
            ])
        charge_output = output.get("charge_output", {})
        charge_evidence = charge_output.get("charge_evidence", [])
        for evidence in charge_evidence:
            sheets["CHARGE_EVIDENCE"].append([
                case_id,
                strategy,
                evidence.get("charge"),
                "; ".join(evidence.get("defendant_entity_ids", [])),
                "; ".join(evidence.get("defendant_names", [])),
                evidence.get("source_region"),
                evidence.get("page_number"),
                "; ".join(evidence.get("line_ids", [])),
                evidence.get("match_method"),
                evidence.get("confidence"),
                evidence.get("raw_text"),
            ])
            sheets["CHARGES"].append([
                case_id,
                evidence.get("charge"),
                evidence.get("source_region"),
                evidence.get("page_number"),
                "; ".join(evidence.get("line_ids", [])),
                evidence.get("match_method"),
                evidence.get("confidence"),
                evidence.get("raw_text"),
            ])
            entity_ids = evidence.get("defendant_entity_ids", [])
            names = evidence.get("defendant_names", [])
            if entity_ids:
                for entity_index, entity_id in enumerate(entity_ids):
                    sheets["DEFENDANT_CHARGES"].append([
                        case_id,
                        entity_id,
                        names[entity_index] if entity_index < len(names) else "",
                        evidence.get("charge"),
                        "mapped",
                        evidence.get("match_method"),
                        "; ".join(evidence.get("line_ids", [])),
                    ])
            else:
                sheets["DEFENDANT_CHARGES"].append([
                    case_id,
                    "",
                    "; ".join(names),
                    evidence.get("charge"),
                    "unmapped",
                    evidence.get("match_method"),
                    "; ".join(evidence.get("line_ids", [])),
                ])
        warning_values = list(charge_output.get("warnings", []))
        warning_values.extend(
            warning
            for warning in output.get("warnings", [])
            if str(warning).startswith((
                "decision_", "charge_", "ambiguous_", "unmatched_", "unresolved_"
            ))
        )
        for warning in dict.fromkeys(str(value) for value in warning_values if value):
            sheets["CHARGE_WARNINGS"].append([case_id, strategy, warning])
        primary_strategy = str(case.get("metrics", {}).get("primary_strategy") or "")
        sentence_output = (
            output.get("sentence_output", {}) if strategy == primary_strategy else {}
        )
        sentence_map = sentence_output.get("defendant_sentence_map", {})
        for defendant in output.get("defendants", []):
            entity_id = str(
                defendant.get("entity_id") or defendant.get("source_block_id") or ""
            )
            sentence = sentence_map.get(entity_id, {}) if isinstance(sentence_map, dict) else {}
            sheets["DEFENDANT_SENTENCES"].append([
                case_id,
                entity_id,
                defendant.get("full_name"),
                sentence.get("primary_penalty_kind"),
                sentence.get("primary_penalty_text"),
                sentence.get("display_text"),
                sentence.get("suspended"),
                sentence.get("probation_text"),
                sentence.get("probation_start_text"),
                sentence.get("execution_status"),
                sentence.get("sentence_start_text"),
                sentence.get("detention_credit_text"),
                sentence.get("completion_text"),
                sentence.get("release_text"),
                sentence.get("aggregate_penalty_text"),
                "; ".join(sentence.get("additional_penalties", [])),
                sentence.get("source_region"),
                sentence.get("page_number"),
                "; ".join(sentence.get("line_ids", [])),
                sentence.get("confidence"),
                "; ".join(sentence.get("warnings", [])),
            ])
        for sentence in sentence_output.get("sentence_evidence", []):
            sheets["SENTENCE_EVIDENCE"].append([
                case_id,
                sentence.get("evidence_type"),
                "; ".join(sentence.get("defendant_entity_ids", [])),
                "; ".join(sentence.get("defendant_names", [])),
                sentence.get("raw_text"),
                sentence.get("page_number"),
                "; ".join(sentence.get("line_ids", [])),
                sentence.get("source_region"),
                sentence.get("match_method"),
                sentence.get("confidence"),
                "; ".join(sentence.get("warnings", [])),
            ])
        for warning in sentence_output.get("sentence_warnings", []):
            sheets["SENTENCE_WARNINGS"].append([
                case_id,
                warning.get("warning"),
                warning.get("severity"),
                warning.get("scope"),
                warning.get("defendant_entity_id"),
                warning.get("defendant_name"),
                warning.get("page_number"),
                "; ".join(warning.get("line_ids", [])),
                warning.get("raw_text"),
                warning.get("source_region"),
                warning.get("front_name"),
                warning.get("decision_name"),
                warning.get("match_method"),
                warning.get("similarity"),
            ])
        for audit in output.get("source_region_audit", []):
            sheets["SOURCE_REGION_AUDIT"].append([
                case_id,
                audit.get("field_name"),
                audit.get("source_region"),
                audit.get("source_page"),
                "; ".join(audit.get("evidence_line_ids", [])),
                audit.get("allowed"),
                audit.get("warning"),
            ])
        sheets["RAW_JSON"].append([case_id, strategy, _display(output)])
        if output.get("needs_review"):
            sheets["NEEDS_REVIEW"].append([case_id, strategy, "; ".join(output.get("warnings", []))])


def _append_trial_panel(sheet, case_id, strategy, panel) -> None:
    roles = [
        ("presiding_judge", panel.get("presiding_judge")),
        *(("juror", name) for name in panel.get("jurors", [])),
        ("clerk", panel.get("clerk")),
        ("prosecutor", panel.get("prosecutor")),
    ]
    for role, name in roles:
        if name:
            sheet.append([case_id, strategy, role, name, None, None, None, None, None])


def _ocr_summary(record: OCRCacheRecord) -> dict[str, Any]:
    metadata = record.result.metadata
    marker = metadata.get("marker", {}) if isinstance(metadata.get("marker"), dict) else {}
    early_stop = metadata.get("early_stop", {}) if isinstance(metadata.get("early_stop"), dict) else {}
    return {
        "source_file": metadata.get("source_file") or record.case_id,
        "ocr_status": record.result.status,
        "marker_found": record.result.marker_found,
        "marker_page": record.result.marker_page or marker.get("page_number"),
        "early_stop_triggered": bool(early_stop.get("triggered")),
        "pages_processed": record.result.pages_processed,
        "pages_total": metadata.get("pages_total"),
    }


def _anchor_llm_unavailable(anchor, case_id, strategy, preflight) -> dict[str, Any]:
    rule = _rule_anchor_output(anchor, case_id, strategy)
    if strategy == "llm_per_block":
        base = empty_pre_content_output(rule["document_type"])
        base["metadata"] = deepcopy(rule["metadata"])
        base["trial_panel"] = deepcopy(rule["trial_panel"])
        base["evidence"] = deepcopy(rule["evidence"])
        base["field_meta"] = deepcopy(rule["field_meta"])
        base["anchor_segments"] = deepcopy(anchor)
        rule = base
    rule.update(
        case_id=case_id,
        strategy=strategy,
        status=f"{strategy}_not_run",
        result_valid=strategy == "rule_then_llm_per_block",
        llm_json_valid=False,
        llm_actually_called=False,
        chunk_count=0,
        llm_status=[_preflight_status(strategy, preflight, required=True)],
        needs_review=True,
    )
    rule["warnings"] = list(rule.get("warnings", [])) + [f"{strategy}_not_run:preflight_failed"]
    return rule


def _hybrid_without_llm(rule, preflight, *, strategy):
    output = deepcopy(rule)
    output.update(
        strategy=strategy,
        status="hybrid_rule_only_fallback",
        result_valid=True,
        llm_json_valid=False,
        llm_actually_called=False,
        chunk_count=0,
        llm_status=[_preflight_status(strategy, preflight, required=False)],
    )
    output["warnings"].append("hybrid_llm_unavailable_rule_output_preserved")
    output["needs_review"] = True
    return output


def _llm_unavailable_output(segment, preflight, *, strategy):
    return {
        "strategy": strategy,
        "document_type": segment.get("document_type"),
        "status": "llm_only_not_run",
        "result_valid": False,
        "llm_json_valid": False,
        "llm_actually_called": False,
        "chunk_count": 0,
        "llm_status": [_preflight_status(strategy, preflight, required=True)],
        "warnings": ["llm_only_not_run:preflight_failed"],
        "needs_review": True,
    }


def _preflight_status(strategy, preflight, *, required):
    value = preflight or {}
    return {
        "strategy": strategy,
        "chunk_name": "preflight",
        "block_type": None,
        "llm_required": required,
        "llm_available": False,
        "llm_actually_called": False,
        "provider": value.get("provider"),
        "model": value.get("model"),
        "base_url": value.get("base_url"),
        "context_window": None,
        "input_chars": 0,
        "estimated_input_tokens": 0,
        "max_output_tokens": None,
        "budget_ok": False,
        "truncated": False,
        "chunked": False,
        "request_ok": False,
        "response_ok": False,
        "error_type": value.get("error_type") or "llm_preflight_failed",
        "error_message": value.get("error") or "Local LLM preflight failed.",
        "duration_ms": 0,
    }


def _attach_decision_tail_charges(
    output: dict[str, Any],
    tail_record: DecisionTailRecord | None,
    *,
    settings: PipelineSettings,
) -> None:
    output.setdefault("warnings", [])
    if tail_record is None:
        output["decision_tail_status"] = "cache_missing"
        output["decision_heading_found"] = False
        output["decision_heading_page"] = None
        output["decision_tail_line_count"] = 0
        output["verdict_candidate_count"] = 0
        output["parsed_charge_count"] = 0
        output["mapped_defendant_count"] = 0
        output["unmapped_verdict_count"] = 0
        output["invalid_charge_count"] = 0
        output["charge_output"] = {
            "case_charges": [],
            "defendant_charge_map": {},
            "charge_evidence": [],
            "verdict_candidate_blocks": [],
            "warnings": ["decision_tail_cache_missing"],
        }
        output["case_charges"] = []
        output["defendant_charge_map"] = {}
        output["sentence_output"] = empty_sentence_output(
            "decision_tail_cache_missing"
        )
        output["defendant_sentence_map"] = {}
        output["sentence_evidence"] = []
        output["sentence_warnings"] = []
        output["warnings"].append("decision_tail_cache_missing")
        output["needs_review"] = True
        return
    defendants = _ensure_front_defendant_entity_ids(output.get("defendants", []))
    output["defendants"] = defendants
    decision_lines = [
        line
        for line in tail_record.lines
        if str(line.get("source_region") or "").strip() in {"", DECISION_TAIL}
    ]
    verdict_blocks = iter_verdict_blocks(decision_lines)
    charge_output = parse_explicit_decision_charges(
        tail_record.lines,
        defendants=defendants,
        source_region=DECISION_TAIL,
        min_name_match_score=settings.decision_name_match_min_score,
        name_match_ambiguity_gap=settings.decision_name_match_ambiguity_gap,
    )
    output["charge_output"] = charge_output
    charge_output["verdict_candidate_blocks"] = verdict_blocks
    output["case_charges"] = list(charge_output["case_charges"])
    output["defendant_charge_map"] = dict(charge_output["defendant_charge_map"])
    sentence_output = (
        parse_defendant_sentences(
            tail_record.lines,
            defendants=defendants,
            source_region=DECISION_TAIL,
            case_id=tail_record.case_id,
            min_name_match_score=settings.decision_name_match_min_score,
            name_match_ambiguity_gap=settings.decision_name_match_ambiguity_gap,
        )
        if tail_record.heading_found
        else empty_sentence_output("decision_tail_heading_not_found")
    )
    output["sentence_output"] = sentence_output
    output["defendant_sentence_map"] = dict(
        sentence_output["defendant_sentence_map"]
    )
    output["sentence_evidence"] = list(sentence_output["sentence_evidence"])
    output["sentence_warnings"] = list(sentence_output.get("sentence_warnings", []))
    output.update(sentence_output.get("diagnostics", {}))
    parsed_charge_count = len(charge_output["charge_evidence"])
    mapped_defendant_count = sum(
        bool(charges) for charges in charge_output["defendant_charge_map"].values()
    )
    unmapped_verdict_count = max(0, len(verdict_blocks) - parsed_charge_count) + sum(
        not item.get("defendant_entity_ids")
        for item in charge_output["charge_evidence"]
    )
    output["decision_heading_found"] = bool(tail_record.heading_found)
    output["decision_heading_page"] = tail_record.heading_page
    output["decision_tail_line_count"] = len(decision_lines)
    output["verdict_candidate_count"] = len(verdict_blocks)
    output["parsed_charge_count"] = parsed_charge_count
    output["mapped_defendant_count"] = mapped_defendant_count
    output["unmapped_verdict_count"] = unmapped_verdict_count
    output["invalid_charge_count"] = int(
        charge_output.get("invalid_charge_count") or 0
    )
    output["decision_tail_status"] = (
        "heading_not_found"
        if not tail_record.heading_found
        else "parsed"
        if charge_output["case_charges"]
        else "charge_not_found"
    )
    diagnostic_warnings: list[str] = []
    if tail_record.heading_found and not verdict_blocks:
        diagnostic_warnings.append(
            "decision_heading_found_but_no_verdict_candidate"
        )
    elif verdict_blocks and not charge_output["case_charges"]:
        diagnostic_warnings.append(
            "verdict_candidates_found_but_no_charge_parsed"
        )
    if parsed_charge_count < len(verdict_blocks):
        diagnostic_warnings.append(
            f"verdict_charge_coverage_incomplete:"
            f"{parsed_charge_count}/{len(verdict_blocks)}"
        )
    if charge_output["case_charges"] and not mapped_defendant_count:
        diagnostic_warnings.append("charges_parsed_but_no_defendant_mapped")
    charge_output["warnings"] = list(
        dict.fromkeys([*charge_output.get("warnings", []), *diagnostic_warnings])
    )
    if charge_output["case_charges"]:
        for item in charge_output["charge_evidence"]:
            line_ids = list(item.get("line_ids", []))
            output.setdefault("evidence", []).append(
                {
                    "field": "metadata.legal_relationship",
                    "value": item.get("charge"),
                    "line_id": line_ids[0] if line_ids else None,
                    "text": item.get("raw_text"),
                    "confidence": 0.99,
                    "source": "decision_tail_charge",
                    "source_region": DECISION_TAIL,
                }
            )
        output.setdefault("field_meta", {})["metadata.legal_relationship"] = {
            "source": "decision_tail_charge",
            "source_region": DECISION_TAIL,
            "confidence": 0.99,
            "evidence_line_ids": list(
                dict.fromkeys(
                    line_id
                    for item in charge_output["charge_evidence"]
                    for line_id in item.get("line_ids", [])
                )
            ),
        }
    else:
        output["warnings"].append("charge_not_found_in_decision_tail")
        output["needs_review"] = True
    output["warnings"].extend(charge_output.get("warnings", []))
    output["warnings"].extend(sentence_output.get("warnings", []))
    output["warnings"].extend(tail_record.warnings)
    output["warnings"] = list(dict.fromkeys(output["warnings"]))
    if any(
        warning in SENTENCE_REVIEW_WARNINGS
        for warning in sentence_output.get("warnings", [])
    ):
        output["needs_review"] = True


def _ensure_front_defendant_entity_ids(values: Any) -> list[dict[str, Any]]:
    defendants = []
    for index, raw in enumerate(values if isinstance(values, list) else [], start=1):
        if not isinstance(raw, dict):
            continue
        item = dict(raw)
        item["entity_id"] = str(
            item.get("entity_id")
            or item.get("source_block_id")
            or f"defendant_{index:03d}"
        )
        item["source_region"] = FRONT_PRE_CONTENT
        defendants.append(item)
    return defendants


def _decision_tail_summary(
    tail_record: DecisionTailRecord | None,
) -> dict[str, Any] | None:
    if tail_record is None:
        return None
    return {
        "status": tail_record.status,
        "heading_found": tail_record.heading_found,
        "heading_page": tail_record.heading_page,
        "pages_scanned": len(tail_record.scanned_page_numbers),
        "warnings": list(tail_record.warnings),
    }


def _skipped_output(segment: dict[str, Any], strategy: str, reason: str) -> dict[str, Any]:
    return {
        "strategy": strategy,
        "document_type": str(segment.get("document_type") or "unknown"),
        "status": "strategy_skipped",
        "result_valid": False,
        "llm_json_valid": False,
        "llm_actually_called": False,
        "chunk_count": 0,
        "llm_status": [],
        "warnings": [reason],
        "needs_review": True,
        "skipped": True,
    }


def _evidence_text(item, line_text):
    return " | ".join(
        line_text.get(str(line_id), "")
        for line_id in item.get("evidence_line_ids", [])
        if line_text.get(str(line_id))
    )


def _write_json(path: Path, payload: Any) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def _display(value: Any) -> str:
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return "" if value is None else str(value)


def _evidence_table(output: dict[str, Any], strategy: str) -> str:
    rows = []
    for item in output.get("evidence", []):
        line_id = str(item.get("line_id") or "")
        link = f'<a href="#{escape(line_id)}">{escape(line_id)}</a>' if line_id else ""
        rows.append(
            f'<tr><td>{escape(strategy)}</td><td>{escape(item.get("field"))}</td>'
            f'<td>{escape(item.get("value"))}</td><td>{link}</td><td>{escape(item.get("text"))}</td></tr>'
        )
    return (
        '<table><tr><th>Strategy</th><th>Field</th><th>Value</th><th>Source line</th><th>Evidence text</th></tr>'
        + "".join(rows) + "</table>"
    )
